"""Worker-Ausgabe gegen einen echten Extrakt abgleichen — zeilengenau je Art.

Prüft vor dem Scharfschalten einer 🟡-Art, ob `worker.py` die Maske wirklich
reproduziert. Beide Dateien liegen im selben Format (der Worker baut genau das
Extrakt-Format), der Vergleich ist also Datenstand-gleich → jede Differenz ist
echte Mapping-/Formel-Differenz.

Ablauf:
    1) Worker-Datei erzeugen (kein Versand):
         py -3-32 worker.py --art wareneingaenge --dry-run
       schreibt odbc-worker/_dryrun_<Dateiname>.
    2) frischen Extrakt derselben Maske ziehen (z. B. nach extrakte/).
    3) abgleichen:
         python worker_abgleich.py wareneingaenge extrakte/AswKpf_WE.txt

Standardmäßig wird die Worker-Datei `_dryrun_<Dateiname>` neben diesem Skript
gesucht; mit --worker <pfad> überschreibbar. Dieses Skript liest nur Dateien,
kein ODBC — läuft mit normalem (64-bit-)Python.

Hinweis: In --dry-run teilen sich einige Arten denselben Dateinamen
(auftraege/auftragspositionen → AswKpf_AUF.txt; wareneingaenge/materialpreise →
AswKpf_WE.txt). Darum je Art einzeln dry-run + abgleichen.
"""
from __future__ import annotations

import argparse
import csv
import io
import sys
from dataclasses import dataclass, field
from pathlib import Path

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

HIER = Path(__file__).resolve().parent


@dataclass(frozen=True)
class Cfg:
    datei: str                       # erwarteter Dateiname (= Worker-Ausgabe)
    kopf: tuple[str, ...]            # Spalten, an denen die Kopfzeile erkannt wird
    key: tuple[str, ...]             # Geschäftsschlüssel (leer = nur Aggregat)
    num: tuple[str, ...] = ()        # numerische Vergleichsspalten (deutsch)
    dat: tuple[str, ...] = ()        # Datums-Vergleichsspalten
    agg: tuple[str, ...] = ()        # Gruppierung für den Aggregat-Modus
    excel: bool = False


# Die sieben 🟡-Arten (die ✅-Arten können hier genauso geprüft werden).
ARTEN: dict[str, Cfg] = {
    "auftragspositionen": Cfg("AswKpf_AUF.txt", ("Vorgang Nr.", "Pos"),
                              ("Vorgang Nr.", "Pos", "UPos"), ("Menge", "Pos Wert"), ("Lieferdatum",)),
    "lieferscheine": Cfg("AswKpf_LS.xlsx", ("Vorgang Nr.", "Pos"),
                         ("Vorgang Nr.", "Pos", "UPos"), ("Menge",), ("Lieferdatum",), excel=True),
    "wareneingaenge": Cfg("AswKpf_WE.txt", ("Vorgang Nr.", "Pos"),
                          ("Vorgang Nr.", "Pos", "UPos"), ("Menge", "Pos Wert")),
    "materialpreise": Cfg("AswKpf_WE.txt", ("Vorgang Nr.", "Pos"),
                          ("Vorgang Nr.", "Pos", "UPos"), ("Menge", "Pos Wert")),
    "liefertreue": Cfg("dev_excel_Liefertreue_Einkauf.txt", ("Auftrag", "Verzug"),
                       ("Auftrag", "Pos", "UPos"), ("Verzug (Tage)", "Menge"), ("geliefert", "Lieferdatum")),
    # Ohne stabilen Schlüssel → Aggregat (Anzahl + Summen je Gruppe).
    "lagerbewegungen": Cfg("AswLagBew.txt", ("Artikelnr", "BuchDatum"), (),
                           ("Bewegungsmenge",), agg=("Artikelnr",)),
    "pruefungen": Cfg("AswQs2151.txt", ("Datum",), (),
                      ("Buchungs-Menge", "Ausschuss-Menge"), agg=("Artikel",)),
}


def _ohne_einfassung(s: str) -> str:
    s = (s or "").strip()
    if s.startswith('="') and s.endswith('"'):
        s = s[2:-1]
    elif s.startswith("=") and len(s) > 1:
        s = s[1:]
    return s.strip().strip('"')


def _num(s: str) -> float | None:
    s = _ohne_einfassung(s).replace(".", "").replace(",", ".")
    if not s:
        return None
    try:
        return float(s)
    except ValueError:
        return None


def _dat(s: str) -> str:
    """Auf dd.mm.yyyy normalisieren (ISO/deutsch/mit Uhrzeit)."""
    s = _ohne_einfassung(s).split(" ")[0]
    for roh, fmt in ((s, "%d.%m.%Y"), (s, "%Y-%m-%d")):
        import datetime as dt
        try:
            return dt.datetime.strptime(roh, fmt).strftime("%d.%m.%Y")
        except ValueError:
            pass
    return s


def _zeilen(pfad: Path, cfg: Cfg) -> list[dict[str, str]]:
    """Datei einlesen; Kopfzeile an den cfg.kopf-Spalten erkennen (überspringt
    Titelzeilen)."""
    if cfg.excel:
        from openpyxl import load_workbook
        ws = load_workbook(pfad, read_only=True, data_only=True).active
        matrix = [["" if z is None else str(z) for z in r]
                  for r in ws.iter_rows(values_only=True)]
    else:
        text = pfad.read_bytes().decode("cp1252", errors="replace")
        matrix = list(csv.reader(io.StringIO(text), delimiter="\t"))
    kopf_idx = next((i for i, r in enumerate(matrix)
                     if all(any(_ohne_einfassung(z).startswith(k) for z in r) for k in cfg.kopf)), None)
    if kopf_idx is None:
        raise SystemExit(f"Kopfzeile mit {cfg.kopf} in {pfad.name} nicht gefunden.")
    hdr = [_ohne_einfassung(z) for z in matrix[kopf_idx]]
    out = []
    for r in matrix[kopf_idx + 1:]:
        if not any(z.strip() for z in r):
            continue
        out.append({hdr[i]: (r[i] if i < len(r) else "") for i in range(len(hdr))})
    return out


def _key(zeile: dict[str, str], cfg: Cfg) -> tuple[str, ...]:
    teile = []
    for k in cfg.key:
        v = _ohne_einfassung(zeile.get(k, ""))
        try:
            v = str(int(float(v.replace(",", "."))))  # 100.0 -> 100, Nummern normieren
        except ValueError:
            pass
        teile.append(v)
    return tuple(teile)


def _index(zeilen: list[dict], cfg: Cfg) -> dict[tuple, dict]:
    idx = {}
    for z in zeilen:
        idx[_key(z, cfg)] = z  # bei Doppeln gewinnt die letzte (wie der Upsert)
    return idx


def _schluessel_abgleich(art: str, worker: list[dict], extrakt: list[dict], cfg: Cfg) -> None:
    w, e = _index(worker, cfg), _index(extrakt, cfg)
    kw, ke = set(w), set(e)
    print(f"[{art}] Worker {len(worker)} Zeilen ({len(kw)} Schlüssel) | Extrakt {len(extrakt)} ({len(ke)})")
    print(f"  nur im Worker:  {len(kw - ke)}")
    print(f"  nur im Extrakt: {len(ke - kw)}")
    gemein = kw & ke
    print(f"  gemeinsam: {len(gemein)}")
    for spalte in cfg.num:
        diff = [k for k in gemein if (a := _num(w[k].get(spalte, ""))) is not None
                and (b := _num(e[k].get(spalte, ""))) is not None and abs(a - b) > 0.01]
        sw = sum(_num(w[k].get(spalte, "")) or 0 for k in gemein)
        se = sum(_num(e[k].get(spalte, "")) or 0 for k in gemein)
        print(f"    {spalte}: {len(diff)} Abweichungen | Summe Worker {sw:,.2f} vs Extrakt {se:,.2f}")
        for k in diff[:5]:
            print(f"      {k}: Worker {w[k].get(spalte)!r} vs Extrakt {e[k].get(spalte)!r}")
    for spalte in cfg.dat:
        diff = [k for k in gemein if _dat(w[k].get(spalte, "")) != _dat(e[k].get(spalte, ""))]
        print(f"    {spalte} (Datum): {len(diff)} Abweichungen")
        for k in diff[:5]:
            print(f"      {k}: Worker {_dat(w[k].get(spalte,''))} vs Extrakt {_dat(e[k].get(spalte,''))}")


def _aggregat_abgleich(art: str, worker: list[dict], extrakt: list[dict], cfg: Cfg) -> None:
    print(f"[{art}] (Aggregat, kein Zeilenschlüssel) Worker {len(worker)} | Extrakt {len(extrakt)} Zeilen")
    for spalte in cfg.num:
        sw = sum(_num(z.get(spalte, "")) or 0 for z in worker)
        se = sum(_num(z.get(spalte, "")) or 0 for z in extrakt)
        d = f"{(sw - se) / se:+.1%}" if se else "—"
        print(f"  Summe {spalte}: Worker {sw:,.2f} vs Extrakt {se:,.2f} ({d})")
    for spalte in cfg.agg:
        gw = {_ohne_einfassung(z.get(spalte, "")) for z in worker}
        ge = {_ohne_einfassung(z.get(spalte, "")) for z in extrakt}
        print(f"  distinct {spalte}: Worker {len(gw)} | Extrakt {len(ge)} | "
              f"nur Worker {len(gw - ge)} | nur Extrakt {len(ge - gw)}")


def main() -> None:
    ap = argparse.ArgumentParser(description="Worker-Ausgabe gegen echten Extrakt abgleichen")
    ap.add_argument("art", choices=list(ARTEN))
    ap.add_argument("extrakt", help="Pfad zum frischen echten Extrakt")
    ap.add_argument("--worker", help="Worker-Datei (Vorgabe: _dryrun_<Dateiname> neben worker.py)")
    a = ap.parse_args()
    cfg = ARTEN[a.art]
    wpfad = Path(a.worker) if a.worker else HIER / f"_dryrun_{cfg.datei}"
    if not wpfad.exists():
        raise SystemExit(f"Worker-Datei fehlt: {wpfad}\n  Erst: py -3-32 worker.py --art {a.art} --dry-run")
    worker = _zeilen(wpfad, cfg)
    extrakt = _zeilen(Path(a.extrakt), cfg)
    (_aggregat_abgleich if not cfg.key else _schluessel_abgleich)(a.art, worker, extrakt, cfg)


if __name__ == "__main__":
    main()
