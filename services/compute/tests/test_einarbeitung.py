"""Der Einarbeitungsbogen — die Mappe, ohne LibreOffice."""
from __future__ import annotations

from datetime import date
from io import BytesIO

from openpyxl import load_workbook

from app.einarbeitung.bogen import DAUER, FORMBLATT, Inhalt, baue_xlsx


def blatt(daten: bytes):
    return load_workbook(BytesIO(daten)).active


def finde(blatt, text: str) -> int | None:
    for zeile in range(1, blatt.max_row + 1):
        for spalte in range(1, 9):
            wert = blatt.cell(zeile, spalte).value
            if wert and text in str(wert):
                return zeile
    return None


INHALTE = [
    Inhalt("Produktion", "Frau Meier", "Sicherheitsunterweisung an der Fräse"),
    Inhalt("QM", "Frau Groß", "Prüfmittel und Prüfpläne"),
]


class TestKopf:
    def test_die_kopfzeilentabelle_traegt_formblatt_und_titel(self):
        b = blatt(baue_xlsx("Dana Neu", "CNC Fräser", date(2026, 9, 5), INHALTE))
        assert b.cell(1, 2).value == FORMBLATT
        assert b.cell(2, 2).value == "Einarbeitungsplan"

    def test_name_stelle_und_beginn_stehen_drin(self):
        b = blatt(baue_xlsx("Dana Neu", "CNC Fräser", date(2026, 9, 5), INHALTE))
        zeile = finde(b, "Name:")
        assert b.cell(zeile, 4).value == "Dana Neu"
        assert b.cell(zeile + 1, 4).value == "CNC Fräser"
        assert b.cell(zeile + 2, 4).value == "05.09.2026"

    def test_das_ende_der_einarbeitung_wird_gerechnet(self):
        beginn = date(2026, 9, 5)
        b = blatt(baue_xlsx("Dana Neu", None, beginn, INHALTE))
        zeile = finde(b, "Einarbeitung bis:")
        assert b.cell(zeile, 4).value == (beginn + DAUER).strftime("%d.%m.%Y")

    def test_ohne_beginn_bleiben_die_felder_leer(self):
        """Ein Bogen ohne Datum ist brauchbar — er wird von Hand ausgefüllt."""
        b = blatt(baue_xlsx("Dana Neu", None, None, INHALTE))
        zeile = finde(b, "Tätigkeitsbeginn:")
        # openpyxl legt eine leere Zelle als `None` ab, nicht als "".
        assert not b.cell(zeile, 4).value


class TestInhalte:
    def test_jeder_inhalt_bekommt_eine_zeile(self):
        b = blatt(baue_xlsx("Dana Neu", None, None, INHALTE))
        assert finde(b, "Sicherheitsunterweisung an der Fräse") is not None
        assert finde(b, "Prüfmittel und Prüfpläne") is not None

    def test_ansprechpartner_und_abteilung_stehen_daneben(self):
        b = blatt(baue_xlsx("Dana Neu", None, None, INHALTE))
        zeile = finde(b, "Sicherheitsunterweisung an der Fräse")
        assert b.cell(zeile, 2).value == "Produktion"
        assert b.cell(zeile, 3).value == "Frau Meier"

    def test_ohne_inhalte_bleibt_der_bogen_lesbar(self):
        b = blatt(baue_xlsx("Dana Neu", None, None, []))
        assert finde(b, "kein Inhalt hinterlegt") is not None

    def test_die_regeln_und_das_ziel_stehen_drauf(self):
        b = blatt(baue_xlsx("Dana Neu", None, None, INHALTE))
        assert finde(b, "Feedbackgespräch") is not None
        assert finde(b, "Ziel der Einarbeitung") is not None

    def test_die_freigabezeile_steht_am_ende(self):
        b = blatt(baue_xlsx("Dana Neu", None, None, INHALTE))
        assert finde(b, "Freigegeben durch:") is not None

    def test_ein_langer_inhalt_bekommt_mehr_hoehe(self):
        """Sonst schneidet der Druck ihn ab."""
        kurz = blatt(baue_xlsx("X", None, None, [Inhalt("A", "B", "Kurz")]))
        lang = blatt(
            baue_xlsx("X", None, None, [Inhalt("A", "B", "Sehr langer Text " * 12)])
        )
        z_kurz = finde(kurz, "Kurz")
        z_lang = finde(lang, "Sehr langer Text")
        assert lang.row_dimensions[z_lang].height > kurz.row_dimensions[z_kurz].height
