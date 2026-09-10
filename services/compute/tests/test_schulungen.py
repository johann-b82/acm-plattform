"""Die Schulungsübersicht: Einlesen und Zuordnen.

Der Parser bekommt eine nachgebaute Übersicht — transponiert, drei Zeilen je
Schulung, Daten in drei Schreibweisen.
"""
from __future__ import annotations

from datetime import date
from io import BytesIO

import pytest
from openpyxl import Workbook

from app.parsing.schulungsuebersicht import (
    als_datum,
    lies_uebersicht,
    turnus_zu_monaten,
)
from app.schulungen.uebernahme import (
    OhneZuordnung,
    personalnummer_aus,
    sammle_ohne_zuordnung,
)


def uebersicht(
    personen=(("101", "Meier, Anna", "NÄH"), ("102", "Klein, Bert", "CUT")),
    schulungen=(
        ("jährlich", "Brandschutz", [(date(2024, 3, 1), date(2025, 3, 1), "Q1/2026"), (None, None, "")]),
        ("bei Bedarf", "Gabelstapler", [(None, None, ""), (date(2023, 5, 1), date(2024, 5, 1), "")]),
    ),
) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Produktion"
    ws.cell(3, 5, "Pers.Nr.")
    ws.cell(4, 5, "Name, Vorname")
    ws.cell(5, 5, "Abt.")
    for i, (nummer, name, abteilung) in enumerate(personen):
        ws.cell(3, 6 + i, nummer)
        ws.cell(4, 6 + i, name)
        ws.cell(5, 6 + i, abteilung)

    zeile = 7
    for turnus, name, werte in schulungen:
        ws.cell(zeile, 2, turnus)
        ws.cell(zeile, 3, name)
        ws.cell(zeile, 4, "Initial")
        ws.cell(zeile + 1, 4, "aktuell")
        ws.cell(zeile + 2, 4, "nächste")
        for i, (initial, aktuell, naechste) in enumerate(werte):
            if initial:
                ws.cell(zeile, 6 + i, initial)
            if aktuell:
                ws.cell(zeile + 1, 6 + i, aktuell)
            if naechste:
                ws.cell(zeile + 2, 6 + i, naechste)
        zeile += 3

    puffer = BytesIO()
    wb.save(puffer)
    return puffer.getvalue()


class TestParser:
    def test_liest_schulungen_und_teilnahmen(self):
        gelesen = lies_uebersicht(uebersicht())
        assert [s.name for s in gelesen.schulungen] == ["Brandschutz", "Gabelstapler"]
        assert gelesen.teilnahmen == 2

    def test_der_bereich_kommt_vom_blattnamen(self):
        assert lies_uebersicht(uebersicht()).schulungen[0].bereich == "Produktion"

    def test_gesamt_wird_aus_dem_bereich_gestrichen(self):
        """Das Blatt heißt „betrieblich (gesamt)", der Bereich „betrieblich"."""
        wb = Workbook()
        ws = wb.active
        ws.title = "betrieblich (gesamt)"
        ws.cell(3, 5, "Pers.Nr.")
        ws.cell(3, 6, "101")
        ws.cell(7, 3, "Brandschutz")
        ws.cell(7, 4, "Initial")
        ws.cell(7, 6, date(2024, 1, 1))
        puffer = BytesIO()
        wb.save(puffer)
        assert lies_uebersicht(puffer.getvalue()).schulungen[0].bereich == "betrieblich"

    def test_zeilen_ohne_aussage_kommen_nicht_mit(self):
        """Die Excel führt alle Personen in allen Zeilen — die meisten Zellen
        sind leer."""
        gelesen = lies_uebersicht(uebersicht())
        brandschutz = gelesen.schulungen[0]
        assert [t.personalnummer for t in brandschutz.teilnahmen] == ["101"]

    def test_abteilungskuerzel_kommt_mit(self):
        gelesen = lies_uebersicht(uebersicht())
        assert gelesen.schulungen[0].teilnahmen[0].abteilung_kuerzel == "NÄH"

    def test_ein_blatt_ohne_kopfzeile_wird_gemeldet(self):
        wb = Workbook()
        wb.active.cell(1, 1, "Nur Text")
        puffer = BytesIO()
        wb.save(puffer)
        gelesen = lies_uebersicht(puffer.getvalue())
        assert gelesen.schulungen == []
        assert "Pers.Nr." in gelesen.hinweise[0]

    def test_unbekannter_turnus_wird_gemeldet(self):
        daten = uebersicht(
            schulungen=(("alle 7 Monde", "Mondschulung", [(date(2024, 1, 1), None, "")]),)
        )
        gelesen = lies_uebersicht(daten)
        assert gelesen.schulungen[0].turnus_monate is None
        assert any("Mondschulung" in h for h in gelesen.hinweise)


class TestTurnus:
    @pytest.mark.parametrize(
        "text, monate",
        [
            ("jährlich", 12),
            ("Jährlich", 12),
            ("alle 2 Jahre", 24),
            ("alle 2 Jahre (und bei Bedarf)", 24),
            # Eine Spanne ergibt keine Frist — geraten wird nicht.
            ("alle 3 - 5 Jahre", None),
            ("bei Bedarf", None),
            (None, None),
            ("Mondphasen", None),
        ],
    )
    def test_uebersetzung(self, text, monate):
        assert turnus_zu_monaten(text) == monate


class TestDatum:
    @pytest.mark.parametrize(
        "wert, erwartet",
        [
            (date(2024, 3, 1), date(2024, 3, 1)),
            # Eine reine Jahreszahl wird der 1. Januar, damit die Reihenfolge stimmt.
            (2024, date(2024, 1, 1)),
            ("2024", date(2024, 1, 1)),
            ("01.03.2024", date(2024, 3, 1)),
            ("2024-03-01", date(2024, 3, 1)),
            ("demnächst", None),
            (None, None),
            ("", None),
        ],
    )
    def test_lesarten(self, wert, erwartet):
        assert als_datum(wert) == erwartet


class TestPersonalnummer:
    def test_wird_ueber_die_beschriftung_gefunden(self):
        """Nicht über die Feld-Kennung: die ist je Personio-Instanz eine andere."""
        roh = {
            "attributes": {
                "dynamic_98765": {"label": "DATEV Personalnummer", "value": "101"},
                "dynamic_11111": {"label": "Kostenstelle", "value": "4711"},
            }
        }
        assert personalnummer_aus(roh) == "101"

    def test_ohne_feld_bleibt_es_leer(self):
        assert personalnummer_aus({"attributes": {}}) is None
        assert personalnummer_aus(None) is None
        assert personalnummer_aus({"attributes": {"x": {"label": "Andere", "value": "1"}}}) is None


class TestOhneZuordnung:
    def test_zeilen_ohne_treffer_werden_gesammelt_nicht_verworfen(self):
        gelesen = lies_uebersicht(uebersicht())
        offen = sammle_ohne_zuordnung(gelesen, {"101": 7})
        assert offen == [
            OhneZuordnung(personalnummer="102", mitarbeiter_name="Klein, Bert", teilnahmen=1)
        ]

    def test_mehrere_teilnahmen_derselben_person_zaehlen_zusammen(self):
        daten = uebersicht(
            personen=(("102", "Klein, Bert", "CUT"),),
            schulungen=(
                ("jährlich", "A", [(date(2024, 1, 1), None, "")]),
                ("jährlich", "B", [(date(2024, 1, 1), None, "")]),
            ),
        )
        offen = sammle_ohne_zuordnung(lies_uebersicht(daten), {})
        assert offen[0].teilnahmen == 2
