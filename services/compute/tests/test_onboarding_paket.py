"""Formblatt 71 und das Onboarding-Paket — die Mappe, ohne LibreOffice.

Das Blatt wird gedruckt und von Hand abgehakt. Geprüft wird deshalb vor allem,
was **nicht** darin steht: für einen Neueintritt bleiben Zeitraum und Kreuze
leer. Ein Datum vorzugeben, das noch niemand terminiert hat, wäre erfunden —
und stünde hinterher gedruckt im Ordner.
"""
from __future__ import annotations

from datetime import date
from io import BytesIO

from openpyxl import load_workbook

from app.einarbeitung.bogen import Inhalt
from app.onboarding.paket import baue_xlsx as paket_xlsx
from app.onboarding.uebersicht import (
    FORMBLATT,
    SP_EX,
    SP_IN,
    SP_JA,
    SP_NEIN,
    TITEL,
    Zeile,
    baue_xlsx,
)

SCHULUNGEN = [
    Zeile(bezeichnung="Produktion: Brandschutzunterweisung"),
    Zeile(bezeichnung="Produktion: Gabelstapler"),
]
EINARBEITUNG = [
    Inhalt("Produktion", "Frau Meier", "Sicherheitsunterweisung an der Fräse"),
]


def blatt(daten: bytes, index: int = 0):
    return load_workbook(BytesIO(daten)).worksheets[index]


def finde(b, text: str) -> int | None:
    for zeile in range(1, b.max_row + 1):
        for spalte in range(1, 8):
            wert = b.cell(zeile, spalte).value
            if wert and text in str(wert):
                return zeile
    return None


class TestKopf:
    def test_formblatt_und_titel(self):
        b = blatt(baue_xlsx("Dana Neu", "CNC Fräser", SCHULUNGEN))
        assert finde(b, FORMBLATT) is not None
        assert finde(b, TITEL) is not None

    def test_name_und_funktion(self):
        b = blatt(baue_xlsx("Dana Neu", "CNC Fräser", SCHULUNGEN))
        zeile = finde(b, "Name:")
        assert b.cell(zeile, 2).value == "Dana Neu"
        assert b.cell(zeile + 1, 2).value == "CNC Fräser"

    def test_ohne_funktion_ein_strich(self):
        """Eine leere Zelle sähe aus wie vergessen."""
        b = blatt(baue_xlsx("Dana Neu", "", SCHULUNGEN))
        assert b.cell(finde(b, "Funktion:"), 2).value == "—"


class TestZeilen:
    def test_laufende_nummer_zweistellig(self):
        b = blatt(baue_xlsx("Dana Neu", "CNC Fräser", SCHULUNGEN))
        erste = finde(b, "Brandschutzunterweisung")
        assert b.cell(erste, 1).value == "01"
        assert b.cell(erste + 1, 1).value == "02"

    def test_fuer_einen_neueintritt_bleibt_alles_offen(self):
        b = blatt(baue_xlsx("Dana Neu", "CNC Fräser", SCHULUNGEN))
        r = finde(b, "Brandschutzunterweisung")
        assert not b.cell(r, 2).value  # Zeitraum
        for spalte in (SP_IN, SP_EX, SP_JA, SP_NEIN):
            assert not b.cell(r, spalte).value

    def test_intern_setzt_genau_ein_kreuz(self):
        b = blatt(baue_xlsx("Dana Neu", "", [Zeile("Brandschutz", intern=True)]))
        r = finde(b, "Brandschutz")
        assert b.cell(r, SP_IN).value == "X"
        assert not b.cell(r, SP_EX).value

    def test_extern_setzt_das_andere(self):
        b = blatt(baue_xlsx("Dana Neu", "", [Zeile("Stapler", intern=False)]))
        r = finde(b, "Stapler")
        assert not b.cell(r, SP_IN).value
        assert b.cell(r, SP_EX).value == "X"

    def test_nachweis_ja_und_nein(self):
        b = blatt(
            baue_xlsx(
                "D",
                "",
                [Zeile("Ersthelfer", nachweis=True), Zeile("Gabelstapler", nachweis=False)],
            )
        )
        assert b.cell(finde(b, "Ersthelfer"), SP_JA).value == "X"
        assert b.cell(finde(b, "Gabelstapler"), SP_NEIN).value == "X"

    def test_anbieter_steht_unter_der_bezeichnung(self):
        """So ist das Formular gebaut — keine eigene Spalte."""
        b = blatt(baue_xlsx("D", "", [Zeile("Stapler", anbieter="TÜV Süd, Würzburg")]))
        text = b.cell(finde(b, "Stapler"), 3).value
        assert text.startswith("Stapler")
        assert "TÜV Süd, Würzburg" in text

    def test_ohne_schulungen_bleibt_der_tabellenkopf(self):
        """Ein leeres Formular ist brauchbar — es wird von Hand ausgefüllt."""
        b = blatt(baue_xlsx("Dana Neu", "CNC Fräser", []))
        assert finde(b, "Bezeichnung der Aus- und Fortbildungsmaßnahme") is not None


class TestSeiteneinrichtung:
    def test_a4_nicht_letter(self):
        """LibreOffices Vorgabe ist Letter; gedruckt wird in Deutschland."""
        b = blatt(baue_xlsx("Dana Neu", "", SCHULUNGEN))
        # openpyxl liest die Größe als Zahl zurück, schreibt sie aber als Text.
        assert int(b.page_setup.paperSize) == int(b.PAPERSIZE_A4)

    def test_tabellenkopf_wiederholt_sich(self):
        b = blatt(baue_xlsx("Dana Neu", "", SCHULUNGEN))
        assert b.print_title_rows

    def test_fuss_ist_eine_echte_fusszeile(self):
        """Als Tabellenzeile stünde er bei wenigen Schulungen mitten auf dem Blatt."""
        b = blatt(baue_xlsx("Dana Neu", "", SCHULUNGEN))
        assert "Ausgabedatum" in b.oddFooter.right.text
        # Kein „\n": openpyxl kodiert es als _x000a_, und LibreOffice gibt das
        # wörtlich aus.
        assert "\n" not in b.oddFooter.left.text


class TestPaket:
    def test_zwei_blaetter_in_einer_mappe(self):
        """Ein Dokument zur Übergabe, nicht zwei — und ohne PDF-Zusammenführung."""
        mappe = load_workbook(
            BytesIO(paket_xlsx("Dana Neu", "CNC Fräser", date(2026, 9, 5),
                               EINARBEITUNG, SCHULUNGEN))
        )
        assert [b.title for b in mappe.worksheets] == [
            "Einarbeitungsplan",
            "Schulungsübersicht",
        ]

    def test_einarbeitung_zuerst(self):
        """Erst die vier Wochen, dann was darüber hinaus ansteht."""
        erstes = blatt(
            paket_xlsx("Dana Neu", "CNC Fräser", date(2026, 9, 5), EINARBEITUNG, SCHULUNGEN),
            0,
        )
        assert finde(erstes, "Sicherheitsunterweisung an der Fräse") is not None

    def test_beide_blaetter_tragen_den_namen(self):
        daten = paket_xlsx("Dana Neu", "CNC Fräser", date(2026, 9, 5), EINARBEITUNG, SCHULUNGEN)
        for index in (0, 1):
            assert finde(blatt(daten, index), "Dana Neu") is not None
