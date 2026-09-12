"""Gemeinsame Containerbeschriftung (ATR-06) gegen eine echte Datenbank.

Wie im Altsystem: die ausgewählten Lieferungen bekommen die Containernummer,
und das Etikett zeigt alles, was diesem Container zugeordnet ist.
"""
from __future__ import annotations

from io import BytesIO

import pytest
import pytest_asyncio
import sqlalchemy as sa
from asgi_lifespan import LifespanManager
from docx import Document
from fastapi import HTTPException
from httpx import ASGITransport, AsyncClient

from app.db import SessionLocal
from app.main import app
from app.routers import atr as atr_router
from tests._auth import mint


@pytest_asyncio.fixture
async def db(datenbank_da):
    if not datenbank_da:
        pytest.skip("Keine Test-Datenbank erreichbar")

    async def leeren():
        async with SessionLocal() as s:
            async with s.begin():
                await s.execute(sa.text("delete from public.atr_lieferungen"))

    await leeren()
    yield
    await leeren()


async def ausfuehren(sql: str, **params):
    async with SessionLocal() as s:
        async with s.begin():
            ergebnis = await s.execute(sa.text(sql), params)
            return [dict(r) for r in ergebnis.mappings()] if ergebnis.returns_rows else []


async def lieferung(ba: str, container: str | None = None) -> str:
    zeilen = await ausfuehren(
        "insert into public.atr_lieferungen (quelle_dateiname, ba_auftrag, containernummer)"
        " values ('ls.pdf', :ba, :c) returning id",
        ba=ba,
        c=container,
    )
    return str(zeilen[0]["id"])


def text_aus(daten: bytes) -> str:
    doc = Document(BytesIO(daten))
    teile = [p.text for p in doc.paragraphs]
    for tabelle in doc.tables:
        for zeile in tabelle.rows:
            for zelle in zeile.cells:
                teile.extend(p.text for p in zelle.paragraphs)
    return "\n".join(teile)


class TestContainerEtikett:
    async def test_weist_zu_und_zeigt_den_ganzen_container(self, db):
        eins = await lieferung("1001")
        zwei = await lieferung("1002")
        schon_drin = await lieferung("1003", "C-9")
        anderswo = await lieferung("1004", "C-1")

        name, daten = await atr_router.container_etikett_bauen(" C-9 ", [eins, zwei])

        assert name == "Container_C-9.docx"
        zeilen = {
            str(z["id"]): z["containernummer"]
            for z in await ausfuehren("select id, containernummer from public.atr_lieferungen")
        }
        assert zeilen[eins] == zeilen[zwei] == zeilen[schon_drin] == "C-9"
        assert zeilen[anderswo] == "C-1"

        text = text_aus(daten)
        assert "Container C-9" in text
        for ba in ("BA 1001", "BA 1002", "BA 1003"):
            assert ba in text
        assert "BA 1004" not in text

    async def test_eine_unbekannte_lieferung_aendert_nichts(self, db):
        eins = await lieferung("1001")
        with pytest.raises(HTTPException) as fehler:
            await atr_router.container_etikett_bauen(
                "C-9", [eins, "00000000-0000-0000-0000-000000000000"]
            )
        assert fehler.value.status_code == 404
        zeile = (await ausfuehren("select containernummer from public.atr_lieferungen"))[0]
        assert zeile["containernummer"] is None

    async def test_ohne_nummer_oder_auswahl_geht_nichts(self, db):
        eins = await lieferung("1001")
        for nummer, auswahl in (("  ", [eins]), ("C-9", [])):
            with pytest.raises(HTTPException) as fehler:
                await atr_router.container_etikett_bauen(nummer, auswahl)
            assert fehler.value.status_code == 422


@pytest_asyncio.fixture
async def client():
    async with LifespanManager(app):
        async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as ac:
            yield ac


class TestRoute:
    async def test_lesen_reicht_nicht(self, db, client):
        eins = await lieferung("1001")
        antwort = await client.post(
            "/api/atr/container-etikett",
            headers={"Authorization": f"Bearer {mint({'atr': 'viewer'})}"},
            json={"containernummer": "C-9", "lieferungen": [eins]},
        )
        assert antwort.status_code == 403

    async def test_liefert_das_word_dokument(self, db, client):
        eins = await lieferung("1001")
        antwort = await client.post(
            "/api/atr/container-etikett",
            headers={"Authorization": f"Bearer {mint({'atr': 'editor'})}"},
            json={"containernummer": "C-9", "lieferungen": [eins]},
        )
        assert antwort.status_code == 200
        assert "Container_C-9.docx" in antwort.headers["content-disposition"]
        assert "BA 1001" in text_aus(antwort.content)
