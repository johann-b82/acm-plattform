"""Der Wartungsbogen — das Raster, ohne LibreOffice.

Geprüft wird die Mappe, nicht das PDF: welche Blätter entstehen, welche
Spalten sie tragen und in welcher Reihenfolge die Aufgaben stehen.
"""
from __future__ import annotations

from datetime import datetime, timezone
from io import BytesIO

import pytest
from openpyxl import load_workbook

from app.wartung.bogen import baue_xlsx, intervall_text

MASCHINE = {
    "name": "Fräse 3",
    "inventarnummer": "INV-0042",
    "standort": "Halle 2",
    "verantwortlich": "Meier",
}


def aufgabe(titel, intervall, wochen=None, minute=0):
    return {
        "titel": titel,
        "intervall": intervall,
        "wochen": wochen,
        "erstellt_am": datetime(2026, 1, 1, 8, minute, tzinfo=timezone.utc),
    }


def blaetter(daten: bytes):
    return load_workbook(BytesIO(daten))


def kopfzeile(blatt, erste_spalte: str) -> int:
    """Die Zeile mit den Spaltenüberschriften — gesucht, nicht geraten."""
    for z in range(1, 20):
        if blatt.cell(z, 1).value == erste_spalte:
            return z
    raise AssertionError(f"Kopfzeile mit {erste_spalte!r} nicht gefunden")


class TestIntervalltext:
    def test_bekannte_intervalle_heissen_deutsch(self):
        assert intervall_text(aufgabe("x", "quartalsweise")) == "Quartalsweise"

    def test_alle_n_wochen_traegt_seine_zahl(self):
        assert intervall_text(aufgabe("x", "alle_n_wochen", 6)) == "Alle 6 Wochen"


class TestBogen:
    def test_erstes_halbjahr_traegt_kw_1_bis_26(self):
        blatt = blaetter(baue_xlsx(MASCHINE, [aufgabe("Ölstand", "monatlich")], 2026, 1))[
            "Periodisch"
        ]
        z = kopfzeile(blatt, "Intervall")
        kopf = [blatt.cell(z, s).value for s in range(3, 3 + 26)]
        assert kopf[0] == "KW 01"
        assert kopf[-1] == "KW 26"

    def test_zweites_halbjahr_traegt_kw_27_bis_52(self):
        blatt = blaetter(baue_xlsx(MASCHINE, [aufgabe("Ölstand", "monatlich")], 2026, 2))[
            "Periodisch"
        ]
        z = kopfzeile(blatt, "Intervall")
        assert blatt.cell(z, 3).value == "KW 27"
        assert blatt.cell(z, 3 + 25).value == "KW 52"

    def test_ohne_taegliche_aufgabe_kein_zweites_blatt(self):
        mappe = blaetter(baue_xlsx(MASCHINE, [aufgabe("Ölstand", "monatlich")], 2026, 1))
        assert mappe.sheetnames == ["Periodisch"]

    def test_mit_taeglicher_aufgabe_ein_tagesblatt(self):
        mappe = blaetter(
            baue_xlsx(MASCHINE, [aufgabe("Sichtprüfung", "taeglich")], 2026, 1)
        )
        assert mappe.sheetnames == ["Periodisch", "Täglich"]
        tag = mappe["Täglich"]
        z = kopfzeile(tag, "Wartungsaufgabe")
        assert tag.cell(z, 2).value == "1"
        assert tag.cell(z, 32).value == "31"

    def test_taegliche_aufgaben_stehen_nicht_auf_dem_wochenbogen(self):
        """Sonst stünde eine tägliche Aufgabe in einem KW-Raster, in dem sie
        26 Häkchen bräuchte."""
        mappe = blaetter(
            baue_xlsx(
                MASCHINE,
                [aufgabe("Sichtprüfung", "taeglich"), aufgabe("Ölstand", "monatlich")],
                2026,
                1,
            )
        )
        blatt = mappe["Periodisch"]
        erste = kopfzeile(blatt, "Intervall") + 1
        titel = [blatt.cell(z, 2).value for z in range(erste, erste + 3)]
        assert "Ölstand" in titel
        assert "Sichtprüfung" not in titel

    def test_gruppen_stehen_in_fester_reihenfolge(self):
        gemischt = [
            aufgabe("Quartal", "quartalsweise", minute=1),
            aufgabe("Woche", "woechentlich", minute=2),
            aufgabe("Sechs Wochen", "alle_n_wochen", 6, minute=3),
            aufgabe("Monat", "monatlich", minute=4),
        ]
        blatt = blaetter(baue_xlsx(MASCHINE, gemischt, 2026, 1))["Periodisch"]
        erste = kopfzeile(blatt, "Intervall") + 1
        titel = [blatt.cell(z, 2).value for z in range(erste, erste + 4)]
        assert titel == ["Woche", "Monat", "Quartal", "Sechs Wochen"]

    def test_ohne_aufgaben_bleibt_der_bogen_lesbar(self):
        blatt = blaetter(baue_xlsx(MASCHINE, [], 2026, 1))["Periodisch"]
        erste = kopfzeile(blatt, "Intervall") + 1
        assert "keine Aufgabe" in str(blatt.cell(erste, 1).value)

    def test_der_kopf_traegt_die_maschine(self):
        blatt = blaetter(baue_xlsx(MASCHINE, [], 2026, 1))["Periodisch"]
        assert blatt.cell(1, 1).value == "Wartungsnachweis"
        assert blatt.cell(3, 2).value == "Fräse 3"
        assert blatt.cell(4, 2).value == "INV-0042"

    def test_fehlende_angaben_werden_zum_strich(self):
        blatt = blaetter(baue_xlsx({"name": "Ohne"}, [], 2026, 1))["Periodisch"]
        assert blatt.cell(4, 2).value == "—"
