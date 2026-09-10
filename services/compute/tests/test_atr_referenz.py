"""Parser der ATR-Referenzmappe.

Die Zelladressen sind aus dem Altprojekt übernommen und an einer echten
Vorlage gewachsen. Diese Prüfungen bauen sie nach — sie belegen, dass der
Parser tut, was er soll, und dass er **laut** scheitert, wenn das Layout nicht
stimmt. Was sie nicht belegen können: dass die Adressen zur Mappe passen, die
der Kunde heute schickt. Dafür braucht es ein echtes Blatt.
"""
from __future__ import annotations

from decimal import Decimal
from io import BytesIO

import pytest
from openpyxl import Workbook

from app.parsing.atr_referenz import MappeUnbrauchbar, lies_referenzmappe


def mappe(
    *,
    kopfzeile_ok: bool = True,
    teile: list[tuple] | None = None,
    zweites_blatt: bool = False,
) -> bytes:
    """Baut eine Mappe im Layout, das der Parser erwartet."""
    wb = Workbook()
    ws = wb.active
    ws.title = "ATR"

    ws["D1"] = "Diehl Aviation"
    ws["D2"] = "A350"
    ws["D3"] = "WP 42"
    ws["D4"] = "PS-1"
    ws["D5"] = "ATP-9"
    ws["D6"] = "LS-7"
    ws["D7"] = "REF-100"
    ws["D8"] = "ACM GmbH"
    ws["G8"] = "KS-3"
    ws["G4"] = "NSCM-55"
    ws["G3"] = "25-11"
    ws["F12"] = "Waage 3"

    ws["C13"] = "Part Number" if kopfzeile_ok else "Irgendwas"
    ws["H13"] = "Weight (kg)"

    zeile = 14
    for eintrag in teile or []:
        art = eintrag[0]
        if art == "kategorie":
            ws[f"A{zeile}"] = eintrag[1]
        elif art == "summe":
            ws[f"F{zeile}"] = "Total weight"
        else:
            _, b, c, d, f, g, h = eintrag
            ws[f"B{zeile}"] = b
            ws[f"C{zeile}"] = c
            ws[f"D{zeile}"] = d
            ws[f"F{zeile}"] = f
            ws[f"G{zeile}"] = g
            ws[f"H{zeile}"] = h
        zeile += 1

    if zweites_blatt:
        wb.create_sheet("Noch eins")

    puffer = BytesIO()
    wb.save(puffer)
    return puffer.getvalue()


TEILE = [
    ("kategorie", "Structure"),
    ("teil", "L-1", "VR-1234-56", "Halter links", "D-2000 Iss B", 2, "1,250"),
    ("teil", "L-2", "VR-1234-57", "Halter rechts", "D-2001 Iss A", 1, 0.75),
    ("kategorie", "If applicable"),
    ("kategorie", "Equipment"),
    ("teil", "L-3", "VR-9999-00", "Blende", "D-3000", 4, None),
]


class TestKopfdaten:
    def test_liest_die_festen_zellen(self):
        m = lies_referenzmappe(mappe(teile=TEILE))
        assert m.kopf["kunde"] == "Diehl Aviation"
        assert m.kopf["programm"] == "A350"
        assert m.kopf["nscm"] == "NSCM-55"
        assert m.kopf["waage"] == "Waage 3"


class TestTeile:
    def test_liest_die_teilezeilen(self):
        m = lies_referenzmappe(mappe(teile=TEILE))
        assert [t.teilenummer for t in m.teile] == [
            "VR-1234-56", "VR-1234-57", "VR-9999-00"]
        assert [t.reihenfolge for t in m.teile] == [1, 2, 3]

    def test_uebernimmt_gewicht_mit_komma(self):
        m = lies_referenzmappe(mappe(teile=TEILE))
        assert m.teile[0].gewicht_kg == Decimal("1.250")
        assert m.teile[1].gewicht_kg == Decimal("0.75")
        assert m.teile[2].gewicht_kg is None

    def test_traegt_die_abschnittsueberschrift_nach(self):
        m = lies_referenzmappe(mappe(teile=TEILE))
        assert m.teile[0].kategorie == "Structure"
        # „If applicable" ist keine Überschrift und darf die davor nicht ersetzen.
        assert m.teile[2].kategorie == "Equipment"

    def test_haelt_am_summenblock_an(self):
        teile = TEILE + [
            ("summe",),
            ("teil", "X", "VR-0000-01", "Danach", "D-9", 1, "1"),
        ]
        m = lies_referenzmappe(mappe(teile=teile))
        assert len(m.teile) == 3

    def test_nur_zeilen_mit_vr_nummer(self):
        teile = [("teil", "L", "ABC-1", "Kein Teil", "D-1", 1, "1")]
        m = lies_referenzmappe(mappe(teile=teile))
        assert m.teile == []
        assert "Keine Teilezeilen gefunden." in m.hinweise

    def test_unlesbares_gewicht_wird_gemeldet_nicht_verschluckt(self):
        teile = [("teil", "L", "VR-1", "Teil", "D-1", 1, "ca. 2 kg")]
        m = lies_referenzmappe(mappe(teile=teile))
        assert m.teile[0].gewicht_kg is None
        assert any("nicht lesbar" in h for h in m.hinweise)


class TestRiegel:
    def test_falsches_layout_scheitert_laut(self):
        """Sonst stuenden Kopfdaten aus den falschen Zellen im Katalog."""
        with pytest.raises(MappeUnbrauchbar, match="Zeile 13"):
            lies_referenzmappe(mappe(kopfzeile_ok=False, teile=TEILE))

    def test_mehrere_sichtbare_blaetter_scheitern(self):
        with pytest.raises(MappeUnbrauchbar, match="sichtbares Blatt"):
            lies_referenzmappe(mappe(teile=TEILE, zweites_blatt=True))

    def test_keine_exceldatei(self):
        with pytest.raises(Exception):
            lies_referenzmappe(b"das ist kein xlsx")
