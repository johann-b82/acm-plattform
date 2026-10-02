"""Den Interview-Parser gegen eine nachgebaute Datei — reine Rechnung.

Prüft die Kernregel (Familien nur zwischen „Produkte" und „Weitere Bereiche"),
die Namensnormalisierung, die Datumswandlung und das Überspringen kopfloser
Blätter.
"""
from __future__ import annotations

from datetime import date
from io import BytesIO

import pytest
from openpyxl import Workbook

from app.parsing.interview import InterviewUnbrauchbar, lies_interviewdatei


def _datei() -> bytes:
    wb = Workbook()
    start = wb.active
    start.title = "Start"
    start.cell(1, 1, "ACM Produktionsfähigkeiten")  # kein „Mitarbeiter"-Kopf

    ws = wb.create_sheet("Näherei")
    kopf = [
        "Mitarbeiter", "Personio Team", "Produkte / Produktfamilien",
        "Maschine einrichten", "Nähen", "Weitere Bereiche",
        "Engpass / Trainingsbedarf", "Validiert durch", "Datum", "Notiz",
    ]
    for i, text in enumerate(kopf, start=1):
        ws.cell(4, i, text)
    # Anna: zwei Stufen, validiert am 18.09.2026 (Excel-Serial 46283).
    ws.cell(5, 1, "Anna Meier"); ws.cell(5, 2, "Sewing"); ws.cell(5, 3, "MSN")
    ws.cell(5, 4, 2); ws.cell(5, 5, 3); ws.cell(5, 6, "Zuschnitt")
    ws.cell(5, 8, "Vera"); ws.cell(5, 9, 46283)
    # John: Doppelleerzeichen im Namen, Stufe 7 (ungültig) wird verworfen.
    ws.cell(6, 1, "John  Adusei Mensah"); ws.cell(6, 2, "Sewing")
    ws.cell(6, 4, 0); ws.cell(6, 5, 7)
    # Beschreibungsblock; „Sonderarbeiten" hat keine Spalte (der #REF!-Fall).
    ws.cell(8, 1, "Aufgabenfamilie"); ws.cell(8, 2, "Enthält aus bestehender Matrix")
    ws.cell(9, 1, "Maschine einrichten"); ws.cell(9, 2, "Nähmaschinen einstellen")
    ws.cell(10, 1, "Nähen"); ws.cell(10, 2, "Nähen Stoff/Leder")
    ws.cell(11, 1, "Sonderarbeiten"); ws.cell(11, 2, "Schäume annähen")

    puffer = BytesIO()
    wb.save(puffer)
    return puffer.getvalue()


def test_ein_blatt_das_kopflose_uebersprungen():
    datei = lies_interviewdatei(_datei(), "interview.xlsx")
    assert [b.bereich for b in datei.blaetter] == ["Näherei"]


def test_familien_nur_zwischen_produkte_und_weitere_bereiche():
    blatt = lies_interviewdatei(_datei(), "x.xlsx").blaetter[0]
    # „Sonderarbeiten" ist beschrieben, hat aber keine Spalte → keine Familie.
    assert [f.name for f in blatt.familien] == ["Maschine einrichten", "Nähen"]
    assert blatt.familien[0].beschreibung == "Nähmaschinen einstellen"


def test_person_stufen_und_felder():
    blatt = lies_interviewdatei(_datei(), "x.xlsx").blaetter[0]
    anna = blatt.personen[0]
    assert anna.name == "Anna Meier"
    assert anna.team == "Sewing"
    assert anna.produkte == "MSN"
    assert anna.weitere_bereiche == "Zuschnitt"
    assert anna.validiert_durch == "Vera"
    assert anna.validiert_am == date(2026, 9, 18)
    assert anna.stufen == {"Maschine einrichten": 2, "Nähen": 3}


def test_doppelleerzeichen_und_ungueltige_stufe():
    blatt = lies_interviewdatei(_datei(), "x.xlsx").blaetter[0]
    john = blatt.personen[1]
    assert john.name == "John Adusei Mensah"
    # Stufe 7 ist ungültig und fällt weg; 0 bleibt.
    assert john.stufen == {"Maschine einrichten": 0}


def test_datei_ohne_bereichsblatt_wird_abgelehnt():
    wb = Workbook()
    wb.active.cell(1, 1, "nur Text")
    puffer = BytesIO()
    wb.save(puffer)
    with pytest.raises(InterviewUnbrauchbar):
        lies_interviewdatei(puffer.getvalue(), "leer.xlsx")
