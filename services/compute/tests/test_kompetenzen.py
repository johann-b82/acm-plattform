"""Die Qualifikationsmatrix: Einlesen und Zuordnen.

Der Parser bekommt eine nachgebaute Bereichsdatei; die Zuordnung der
Spaltenköpfe zu Personio wird ohne Datenbank geprüft — sie ist reine Rechnung.
"""
from __future__ import annotations

from datetime import date
from io import BytesIO

import pytest
from openpyxl import Workbook

from app.kompetenzen.uebernahme import baue_index, finde_person, normalisiere
from app.parsing.kompetenzmatrix import MatrixUnbrauchbar, lies_matrixdatei


class Zeile:
    def __init__(self, id_, vorname, nachname):
        self.id = id_
        self.first_name = vorname
        self.last_name = nachname


def mappe(kopfzeile: int = 6, personen=("Anna Meier", "Bert Klein"), mit_stand=True) -> bytes:
    """Eine Bereichsdatei nachbauen, so wie die echten aufgebaut sind."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Produktion"
    ws.cell(1, 2, "Qualifikationsmatrix - Produktion")
    if mit_stand:
        ws.cell(2, 6, "Stand")
        ws.cell(2, 7, date(2026, 3, 1))
    ws.cell(kopfzeile, 6, "Anzahl Mitarbeiter")
    ws.cell(kopfzeile, 7, "Durchschnitt")
    for i, name in enumerate(personen):
        ws.cell(kopfzeile, 9 + i * 2, name)

    zeilen = [
        (1, "Allgemein", "Deutsch", [(4, 100), (2, 50)]),
        (2, "Allgemein", "Englisch", [(None, None), (3, 75)]),
        (3, "Maschinen", "Fräse", [(4, 80), (None, None)]),
    ]
    for versatz, (nr, kategorie, bezeichnung, zellen) in enumerate(zeilen):
        z = kopfzeile + 1 + versatz
        ws.cell(z, 2, nr)
        ws.cell(z, 3, kategorie)
        ws.cell(z, 4, bezeichnung)
        for i, (level, grad) in enumerate(zellen):
            if level is not None:
                ws.cell(z, 9 + i * 2, level)
            if grad is not None:
                ws.cell(z, 10 + i * 2, grad)

    puffer = BytesIO()
    wb.save(puffer)
    return puffer.getvalue()


class TestParser:
    def test_liest_qualifikationen_und_personen(self):
        datei = lies_matrixdatei(mappe(), "produktion.xlsx")
        matrix = datei.matrizen[0]
        assert matrix.personen == ["Anna Meier", "Bert Klein"]
        assert [q.bezeichnung for q in matrix.qualifikationen] == [
            "Deutsch",
            "Englisch",
            "Fräse",
        ]

    @pytest.mark.parametrize("kopf", [4, 5, 6])
    def test_die_kopfzeile_darf_wandern(self, kopf):
        """Die vier Bereichsdateien haben sie auf Zeile 4 bis 6."""
        datei = lies_matrixdatei(mappe(kopfzeile=kopf), "x.xlsx")
        assert len(datei.matrizen[0].qualifikationen) == 3

    def test_leere_zellen_werden_nicht_angelegt(self):
        matrix = lies_matrixdatei(mappe(), "x.xlsx").matrizen[0]
        englisch = matrix.qualifikationen[1]
        assert [b.person for b in englisch.bewertungen] == [1]

    def test_stand_und_titel_kommen_mit(self):
        matrix = lies_matrixdatei(mappe(), "x.xlsx").matrizen[0]
        assert matrix.stand == date(2026, 3, 1)
        assert matrix.titel == "Qualifikationsmatrix - Produktion"

    def test_ohne_stand_bleibt_es_leer(self):
        matrix = lies_matrixdatei(mappe(mit_stand=False), "x.xlsx").matrizen[0]
        assert matrix.stand is None

    def test_werte_ausserhalb_des_bereichs_zaehlen_nicht(self):
        """„N/A" kommt als Text vor, und ein Level 7 gibt es nicht."""
        wb = Workbook()
        ws = wb.active
        ws.cell(6, 6, "Anzahl Mitarbeiter")
        ws.cell(6, 9, "Anna Meier")
        ws.cell(7, 4, "Deutsch")
        ws.cell(7, 9, 7)        # Level ausserhalb 0..4
        ws.cell(7, 10, "N/A")   # Text statt Prozent
        puffer = BytesIO()
        wb.save(puffer)
        matrix = lies_matrixdatei(puffer.getvalue(), "x.xlsx").matrizen[0]
        assert matrix.qualifikationen[0].bewertungen == []

    def test_eine_datei_ohne_matrix_wird_abgelehnt(self):
        wb = Workbook()
        wb.active.cell(1, 1, "Nur Text")
        puffer = BytesIO()
        wb.save(puffer)
        with pytest.raises(MatrixUnbrauchbar):
            lies_matrixdatei(puffer.getvalue(), "leer.xlsx")

    def test_das_diagrammblatt_wird_uebersprungen(self):
        wb = Workbook()
        ws = wb.active
        ws.title = "Qualifikationsdiagramm"
        ws.cell(1, 1, "nur Formeln")
        zweites = wb.create_sheet("Produktion")
        zweites.cell(6, 6, "Anzahl Mitarbeiter")
        zweites.cell(6, 9, "Anna Meier")
        zweites.cell(7, 4, "Deutsch")
        puffer = BytesIO()
        wb.save(puffer)
        datei = lies_matrixdatei(puffer.getvalue(), "x.xlsx")
        assert [m.blatt for m in datei.matrizen] == ["Produktion"]


class TestZuordnung:
    """Eine falsche Zuordnung wäre schlimmer als keine — an ihr hängen
    Leistungsbewertungen."""

    def setup_method(self):
        self.exakt, self.teile = baue_index(
            [
                Zeile(1, "Anna", "Meier"),
                Zeile(2, "Fernando", "Gomes Ferreira"),
                Zeile(3, "Peter", "Meier"),
            ]
        )

    def test_genauer_name(self):
        assert finde_person("Anna Meier", self.exakt, self.teile) == 1

    def test_schreibweise_ist_egal(self):
        assert finde_person("  anna   MEIER ", self.exakt, self.teile) == 1

    def test_fehlender_namensteil_wird_gefunden(self):
        assert finde_person("Fernando Gomes", self.exakt, self.teile) == 2

    def test_mehrdeutig_heisst_keine_zuordnung(self):
        assert finde_person("Meier", self.exakt, self.teile) is None

    def test_unbekannter_name_bleibt_offen(self):
        assert finde_person("Erika Mustermann", self.exakt, self.teile) is None

    def test_ein_teil_der_nirgends_passt_verhindert_den_treffer(self):
        assert finde_person("Anna Meier Schmidt", self.exakt, self.teile) is None

    def test_normalisierung(self):
        assert normalisiere("  Anna   Meier ") == "anna meier"
