"""Die ATR-Erzeugung.

Die Gerüste im Test sind nachgebaut, aber nicht geraten: ihr Aufbau ist an den
**echten** Vorlagen aus der Produktion abgelesen — die A350-Mappe hat die
Tabellenüberschrift in Zeile 13 und die Summenzeile in Zeile 80, die
A380-Mappe in Zeile 10 und 14. Genau diese Verschiebung ist der Grund, warum
Zeilen über Beschriftungen gefunden werden; beide Fälle stehen deshalb hier.

Die echten Vorlagen selbst liegen nicht im Repository: sie tragen
Kundenspezifikationen, Teilenummern und ein Firmenlogo.
"""
from __future__ import annotations

from datetime import date
from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook

from app.atr.etikett import baue_etikett, spaltenzahl
from app.atr.excel import VorlageUnbrauchbar, baue_atr
from app.atr.format import bestellposition, dokumentnummer, programmfamilie

KOPFSPALTEN = [
    "PO Pos", "Supplier Article Code", "Part Number / Index", "Part Name",
    "Serial Number", "Drawing Number / Issue", "Qty", "Weight [kg]",
    "Dimension", "Visual\nInspection", "Material", "Documents",
    "Identification / Markings", "Incpected",
]


def geruest(kopfzeile: int, beispielzeilen: int = 2) -> bytes:
    """Ein Gerüst im Aufbau der echten Vorlagen.

    `kopfzeile` ist die Zeile mit „PO Pos"; darunter Beispielzeilen, dann die
    Summenzeile, dann der Zertifizierungsblock — wie in den Originalen.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "CCRC 6 BED"

    ws["A1"], ws["D1"] = "Customer:", "Diehl Aviation"
    ws["F1"], ws["G1"] = "Purchase Order No.:", "4500000"
    ws["I1"], ws["J1"], ws["K1"] = " 79 - ", "0000", "- 94"
    ws["A2"], ws["D2"] = "AC Programme:", "A350 XWB"
    ws["F2"], ws["G2"] = "MSN:", ""
    ws[f"A{kopfzeile - 4}"] = "Supplier:"
    ws[f"A{kopfzeile - 3}"] = "Manufacturing Process Reference. WO No.: "
    ws[f"A{kopfzeile - 2}"] = "SET PLATZHALTER"
    ws[f"A{kopfzeile - 1}"] = "Weighing date:"
    ws[f"E{kopfzeile - 1}"] = "Weighing Equipment:"
    ws[f"I{kopfzeile - 1}"] = "Testing Date:"

    for spalte, text in enumerate(KOPFSPALTEN, start=1):
        ws.cell(kopfzeile, spalte, text)

    zeile = kopfzeile + 1
    ws.cell(zeile, 1, "BEISPIELABSCHNITT")  # Abschnittsüberschrift
    zeile += 1
    for nummer in range(beispielzeilen):
        ws.cell(zeile, 2, 1000 + nummer)
        ws.cell(zeile, 3, "VR-BEISPIEL")
        ws.cell(zeile, 7, 1)
        ws.cell(zeile, 8, 1.0)
        zeile += 1

    ws.cell(zeile, 6, "Total weight")
    ws.cell(zeile, 8, 0)
    zeile += 1
    ws.cell(zeile, 6, "Max. Guaranteed weight")
    zeile += 2
    ws.cell(zeile, 1, "We hereby certify that the items listed above …")
    zeile += 1
    ws.cell(zeile, 1, "Date:")
    ws.cell(zeile, 2, "=TODAY()")
    ws.cell(zeile, 3, "Ein Name aus der Vorlage")
    ws.cell(zeile, 7, "Date:")
    ws.cell(zeile, 12, "Ein Name aus der Vorlage")

    puffer = BytesIO()
    wb.save(puffer)
    return puffer.getvalue()


LIEFERUNG = {
    "programm": "A350",
    "ba_auftrag": "1021999",
    "bestellnummer": "4500123",
    "msn": "815",
    "satz_titel": "SET 2 BED CCRC",
    "atr_nummer": "4820",
    "wiegedatum": date(2026, 8, 14),
    "pruefdatum": date(2026, 8, 15),
    "qs_unterschrift": "M. Berger",
    "max_gewicht_kg": 12.5,
}

POSITIONEN = [
    {
        "kategorie": "SEC. LINING", "bestellposition": "12",
        "lieferantennummer": "6008", "teilenummer": "VR11S 1020 000 000",
        "bezeichnung": "SEC LINING BED 03", "zeichnung": "VR11S 1020-02/A",
        "menge": 2, "gewicht_kg": "0.679",
        "seriennummern": ["A08UAEL3376", "A08UAEL3377"], "teil_id": "x",
    },
    {
        "kategorie": "SEC. LINING", "bestellposition": "3",
        "lieferantennummer": "6003", "teilenummer": "VR11S 1020 001 000",
        "bezeichnung": "SEC LINING BED 01", "zeichnung": "VR11S 1020-02/A",
        "menge": 1, "gewicht_kg": "0.682", "seriennummern": [], "teil_id": "x",
    },
    {
        "kategorie": "CURTAIN", "bestellposition": "5",
        "lieferantennummer": "6037", "teilenummer": "VR11S 1030 006 000",
        "bezeichnung": "CURTAIN BED 05", "zeichnung": None,
        "menge": 1, "gewicht_kg": None, "seriennummern": [], "teil_id": None,
    },
]


def blatt(daten: bytes):
    return [w for w in load_workbook(BytesIO(daten)).worksheets
            if w.sheet_state == "visible"][0]


def finde(ws, spalte: int, text: str) -> int | None:
    for zeile in range(1, ws.max_row + 1):
        wert = ws.cell(zeile, spalte).value
        if wert and str(wert).strip().lower().startswith(text.lower()):
            return zeile
    return None


class TestBeideLayouts:
    """Der Grund für die Suche über Beschriftungen: die echten Vorlagen haben
    die Tabelle an verschiedenen Stellen."""

    @pytest.mark.parametrize("kopfzeile", [13, 10])
    def test_positionen_landen_unter_der_ueberschrift(self, kopfzeile):
        ws = blatt(baue_atr(geruest(kopfzeile), LIEFERUNG, POSITIONEN))
        ueberschrift = finde(ws, 1, "PO Pos")
        assert ueberschrift == kopfzeile
        # Abschnitt, zwei Teile, Abschnitt, ein Teil.
        assert ws.cell(kopfzeile + 1, 1).value == "SEC. LINING"
        assert ws.cell(kopfzeile + 2, 3).value == "VR11S 1020 000 000"
        assert ws.cell(kopfzeile + 4, 1).value == "CURTAIN"

    @pytest.mark.parametrize("kopfzeile", [13, 10])
    def test_der_zertifizierungsblock_rutscht_mit(self, kopfzeile):
        ws = blatt(baue_atr(geruest(kopfzeile), LIEFERUNG, POSITIONEN))
        summe = finde(ws, 6, "Total weight")
        assert summe is not None
        assert finde(ws, 1, "We hereby certify") > summe


class TestInhalt:
    def test_gewicht_rechnet_mit_der_menge(self):
        ws = blatt(baue_atr(geruest(13), LIEFERUNG, POSITIONEN))
        assert ws.cell(15, 8).value == pytest.approx(1.358)  # 0,679 × 2
        assert ws.cell(16, 8).value == pytest.approx(0.682)

    def test_die_summe_zaehlt_nur_gewogene_zeilen(self):
        ws = blatt(baue_atr(geruest(13), LIEFERUNG, POSITIONEN))
        summe = finde(ws, 6, "Total weight")
        assert ws.cell(summe, 8).value == pytest.approx(2.04)

    def test_hoechstgewicht_landet_darunter(self):
        ws = blatt(baue_atr(geruest(13), LIEFERUNG, POSITIONEN))
        zeile = finde(ws, 6, "Max. Guaranteed weight")
        assert ws.cell(zeile, 8).value == pytest.approx(12.5)
        assert ws.cell(zeile, 8).number_format == "[$-407]0.00"

    def test_hoechstgewicht_aus_der_vorlage_wird_umformatiert(self):
        """Sonst steht auf demselben Blatt „211.00" neben „4,63" — an einem
        echten Dokument aufgefallen."""
        with_wert = geruest(13)
        from openpyxl import load_workbook as _laden

        # Die Vorlage bringt einen eigenen Wert mit, die Lieferung keinen.
        mappe = _laden(BytesIO(with_wert))
        ws0 = mappe.active
        zeile = next(
            z for z in range(1, ws0.max_row + 1)
            if str(ws0.cell(z, 6).value or "").startswith("Max.")
        )
        ws0.cell(zeile, 8, 211.0)
        puffer = BytesIO()
        mappe.save(puffer)

        ws = blatt(baue_atr(puffer.getvalue(),
                            {**LIEFERUNG, "max_gewicht_kg": None}, POSITIONEN))
        z = finde(ws, 6, "Max. Guaranteed weight")
        assert ws.cell(z, 8).value == pytest.approx(211.0)  # Wert bleibt
        assert ws.cell(z, 8).number_format == "[$-407]0.00"  # Format angeglichen

    def test_bestellposition_wird_normiert(self):
        ws = blatt(baue_atr(geruest(13), LIEFERUNG, POSITIONEN))
        assert [ws.cell(z, 1).value for z in (15, 16, 18)] == ["120", "030", "050"]

    def test_seriennummern_sonst_n_a(self):
        ws = blatt(baue_atr(geruest(13), LIEFERUNG, POSITIONEN))
        assert ws.cell(15, 5).value == "A08UAEL3376, A08UAEL3377"
        assert ws.cell(16, 5).value == "N/A"

    def test_kopfdaten_der_lieferung(self):
        ws = blatt(baue_atr(geruest(13), LIEFERUNG, POSITIONEN))
        assert ws.cell(finde(ws, 1, "Manufacturing Process"), 4).value == "1021999"
        assert ws.cell(1, 7).value == "4500123"
        assert ws.cell(2, 7).value == "815"
        # Die A350-Bestellzeile trägt die MSN links aufgefüllt.
        assert ws.cell(1, 10).value == "0815"

    def test_satzbezeichnung_steht_ueber_der_wiegezeile(self):
        ws = blatt(baue_atr(geruest(13), LIEFERUNG, POSITIONEN))
        wiegen = finde(ws, 1, "Weighing date")
        assert ws.cell(wiegen - 1, 1).value == "SET 2 BED CCRC"

    def test_daten_deutsch(self):
        ws = blatt(baue_atr(geruest(13), LIEFERUNG, POSITIONEN))
        wiegen = finde(ws, 1, "Weighing date")
        assert ws.cell(wiegen, 3).value == "14.08.2026"
        assert ws.cell(wiegen, 12).value == "15.08.2026"

    def test_ohne_datum_steht_heute_da_und_nicht_die_formel(self):
        """Die Vorlage trägt `=TODAY()`; LibreOffice setzt das in seiner
        eigenen Sprache — im PDF stand „9/10/2026". Nachgemessen."""
        from datetime import date as _date

        ohne = blatt(baue_atr(
            geruest(13),
            {**LIEFERUNG, "wiegedatum": None, "pruefdatum": None},
            POSITIONEN,
        ))
        wiegen = finde(ohne, 1, "Weighing date")
        heute = _date.today().strftime("%d.%m.%Y")
        assert ohne.cell(wiegen, 3).value == heute
        assert ohne.cell(wiegen, 12).value == heute

    def test_qs_unterschrift_schlaegt_die_vorlage(self):
        """Die Vorlage bringt einen festen Namen mit; maßgeblich ist die
        Lieferung — auch wenn dort nichts steht."""
        ws = blatt(baue_atr(geruest(13), LIEFERUNG, POSITIONEN))
        zeile = finde(ws, 1, "Date:")
        assert ws.cell(zeile, 3).value == "M. Berger"
        assert ws.cell(zeile, 2).value == "15.08.2026"  # kein =TODAY() mehr

        ohne = blatt(baue_atr(geruest(13), {**LIEFERUNG, "qs_unterschrift": None},
                              POSITIONEN))
        # openpyxl legt eine leere Zeichenkette als `None` ab; der feste Name
        # aus der Vorlage ist jedenfalls weg.
        assert not ohne.cell(finde(ohne, 1, "Date:"), 3).value

    def test_a380_ohne_msn(self):
        ws = blatt(baue_atr(geruest(10), {**LIEFERUNG, "programm": "A380", "msn": None},
                            POSITIONEN))
        assert ws.cell(2, 7).value == "N/A Spare Part"


class TestDruck:
    def test_druckkopfzeile_traegt_die_dokumentnummer(self):
        ws = blatt(baue_atr(geruest(13), LIEFERUNG, POSITIONEN))
        assert "ACM-A350CRC-ATR-4820-01" in ws.oddHeader.right.text

    def test_umbrueche_der_kopfzeile_sind_echt(self):
        """openpyxl schreibt `\\n` als `_x000a_`; LibreOffice zeigt das wörtlich.
        Nachgemessen, bevor die Nachbehandlung eingebaut war."""
        import zipfile

        daten = baue_atr(geruest(13), LIEFERUNG, POSITIONEN)
        with zipfile.ZipFile(BytesIO(daten)) as z:
            blatt_xml = next(
                z.read(n).decode("utf-8")
                for n in z.namelist()
                if n.startswith("xl/worksheets/sheet") and n.endswith(".xml")
            )
        kopf = blatt_xml[blatt_xml.find("<headerFooter") : blatt_xml.find("</headerFooter>")]
        assert "_x000a_" not in kopf
        assert "&#10;" in kopf

    def test_auf_eine_seite_breite(self):
        ws = blatt(baue_atr(geruest(13), LIEFERUNG, POSITIONEN))
        assert ws.page_setup.fitToWidth == 1
        assert ws.page_setup.fitToHeight == 0
        assert ws.page_margins.right == ws.page_margins.left


class TestRiegel:
    def test_ohne_tabellenueberschrift_scheitert_es_laut(self):
        wb = Workbook()
        puffer = BytesIO()
        wb.save(puffer)
        with pytest.raises(VorlageUnbrauchbar, match="PO Pos"):
            baue_atr(puffer.getvalue(), LIEFERUNG, POSITIONEN)

    def test_ohne_summenzeile_ebenso(self):
        wb = Workbook()
        ws = wb.active
        for spalte, text in enumerate(KOPFSPALTEN, start=1):
            ws.cell(13, spalte, text)
        puffer = BytesIO()
        wb.save(puffer)
        with pytest.raises(VorlageUnbrauchbar, match="Total weight"):
            baue_atr(puffer.getvalue(), LIEFERUNG, POSITIONEN)


class TestEtikett:
    def test_ein_word_dokument_entsteht(self):
        daten = baue_etikett({**LIEFERUNG, "containernummer": "C-4711"}, POSITIONEN)
        assert daten[:2] == b"PK"  # docx ist ein Zip
        assert len(daten) > 5000

    def test_spaltenzahl_waechst_erst_ab_fuenf(self):
        assert spaltenzahl(1) == 1
        assert spaltenzahl(4) == 1
        assert spaltenzahl(5) == 2
        assert spaltenzahl(8) == 2
        assert spaltenzahl(9) == 3
        assert spaltenzahl(99) == 3  # nie mehr als drei


class TestFormat:
    @pytest.mark.parametrize(
        "roh, erwartet",
        [("1", "010"), ("12", "120"), ("5", "050"), ("120", "120"),
         ("", ""), ("A1", "A1"), (None, None)],
    )
    def test_bestellposition(self, roh, erwartet):
        assert bestellposition(roh) == erwartet

    def test_bestellposition_ist_wiederholbar(self):
        """Sie läuft beim Schreiben und beim Erzeugen."""
        assert bestellposition(bestellposition("12")) == "120"

    @pytest.mark.parametrize(
        "roh, erwartet",
        [
            ("A350 XWB ", "A350"),
            ("A380 - 800", "A380"),
            ("A350", "A350"),
            ("a380", "A380"),
            ("", None),
            (None, None),
            ("Sonstiges", "Sonstiges"),
        ],
    )
    def test_programmfamilie(self, roh, erwartet):
        """Die Mappe schreibt „A350 XWB", der Lieferschein „A350" — an den
        echten Vorlagen nachgesehen."""
        assert programmfamilie(roh) == erwartet

    def test_dokumentnummer(self):
        assert dokumentnummer("4820", "A350") == "ACM-A350CRC-ATR-4820-01 / Issue: 01"
        assert dokumentnummer("4820", "A380") == "ACM-A380-ATR-4820-01 / Issue: 01"
        assert "000" in dokumentnummer(None, "A350")
