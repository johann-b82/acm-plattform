"""Containerbeschriftung als Word-Dokument.

Was auf den Container geklebt wird, damit im Wareneingang klar ist, was
darin liegt: BA-Auftrag groß in Schwarz, Bestellung und Positionen in Rot,
Programm und MSN in Grün.

Hängt an keiner Vorlage — das Dokument entsteht von Grund auf.

Aus `lumeapps` übernommen (`backend/app/services/atr_generate_docx.py`),
unverändert im Verhalten.
"""
from __future__ import annotations

from io import BytesIO
from math import ceil

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor

from app.atr.format import bestellposition

SCHWARZ = RGBColor(0x00, 0x00, 0x00)
ROT = RGBColor(0xC0, 0x00, 0x00)
GRUEN = RGBColor(0x00, 0x80, 0x00)
GROSS = Pt(32)  # BA- und MSN-Zeile — hervorgehoben
NORMAL = Pt(20)

#: Mehrere Lieferungen auf einem Etikett stehen in einem randlosen Raster. Die
#: Schrift bleibt je Spaltenzahl fest — lesbar, nie kleingerechnet; mehr
#: Lieferungen bringen Spalten, darüber hinaus geht es auf Seite zwei weiter.
ZEILEN_JE_SEITE = 4
MAX_SPALTEN = 3
RASTER_SCHRIFT = {1: (Pt(24), Pt(15)), 2: (Pt(20), Pt(13)), 3: (Pt(18), Pt(12))}
RASTER_RAND = Cm(1.5)
RASTER_ABSTAND = Pt(10)


def _zeilen(lieferung: dict, positionen: list[dict], gross=GROSS, normal=NORMAL):
    """(Text, Farbe, Größe) je Zeile."""
    pos = sorted(
        {
            p
            for eintrag in positionen
            if (p := (bestellposition(eintrag.get("bestellposition")) or "").strip())
        }
    )
    # Alle tatsächlich gelieferten Warengruppen in der Reihenfolge des ersten
    # Auftretens — dieselbe Gruppierung wie im ATR.
    gruppen = ", ".join(
        dict.fromkeys(
            k
            for eintrag in positionen
            if (k := (eintrag.get("kategorie") or "").strip())
        )
    )
    return [
        (f"BA {lieferung.get('ba_auftrag') or ''}".rstrip(), SCHWARZ, gross),
        (f"PO {lieferung.get('bestellnummer') or ''}".rstrip(), ROT, normal),
        (f"Pos. {', '.join(pos)}", ROT, normal),
        (
            " ".join(
                s
                for s in (
                    lieferung.get("programm") or "",
                    gruppen,
                    f"MSN {lieferung.get('msn') or ''}".rstrip(),
                )
                if s
            ),
            GRUEN,
            gross,
        ),
    ]


def _querformat(rand=None):
    doc = Document()
    abschnitt = doc.sections[0]
    abschnitt.orientation = WD_ORIENT.LANDSCAPE
    abschnitt.page_width, abschnitt.page_height = (
        abschnitt.page_height,
        abschnitt.page_width,
    )
    if rand is not None:
        abschnitt.left_margin = abschnitt.right_margin = rand
        abschnitt.top_margin = abschnitt.bottom_margin = rand
    return doc


def _setze(absatz, text, farbe, groesse, eng=False):
    absatz.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if eng:
        absatz.paragraph_format.space_before = Pt(0)
        absatz.paragraph_format.space_after = Pt(0)
    lauf = absatz.add_run(text)
    lauf.bold = True
    lauf.font.size = groesse
    lauf.font.color.rgb = farbe
    return absatz


def _als_bytes(doc) -> bytes:
    puffer = BytesIO()
    doc.save(puffer)
    return puffer.getvalue()


def _zeichne(zeilen) -> bytes:
    doc = _querformat()
    for text, farbe, groesse in zeilen:
        _setze(doc.add_paragraph(), text, farbe, groesse)
    return _als_bytes(doc)


def baue_etikett(lieferung: dict, positionen: list[dict]) -> bytes:
    """Etikett für eine Lieferung."""
    zeilen = _zeilen(lieferung, positionen) + [
        (f"Container {lieferung.get('containernummer') or ''}".rstrip(), SCHWARZ, NORMAL),
    ]
    return _zeichne(zeilen)


def spaltenzahl(anzahl: int) -> int:
    """Spalten für `anzahl` Blöcke: eine, solange sie auf eine Seite passen,
    dann zwei oder drei — mehr nie."""
    return max(1, min(MAX_SPALTEN, ceil(anzahl / ZEILEN_JE_SEITE)))


def baue_container_etikett(containernummer: str, lieferungen: list[tuple[dict, list[dict]]]) -> bytes:
    """Ein Etikett für einen ganzen Container.

    Eine oder zwei Lieferungen behalten das große einspaltige Bild. Ab drei
    gehen die Blöcke in ein randloses Raster mit fester, lesbarer Schrift: eine
    Spalte bis vier Blöcke, zwei bis acht, drei darüber — mehr als zwölf laufen
    einfach auf der nächsten Seite weiter.
    """
    anzahl = len(lieferungen)
    if anzahl <= 2:
        zeilen = [(f"Container {containernummer}", SCHWARZ, GROSS)]
        for lieferung, positionen in lieferungen:
            zeilen.append(("", SCHWARZ, Pt(12)))
            zeilen.extend(_zeilen(lieferung, positionen))
        return _zeichne(zeilen)

    doc = _querformat(rand=RASTER_RAND)
    _setze(doc.add_paragraph(), f"Container {containernummer}", SCHWARZ, GROSS)

    spalten = spaltenzahl(anzahl)
    zeilenzahl = ceil(anzahl / spalten)
    gross, normal = RASTER_SCHRIFT[spalten]

    tabelle = doc.add_table(rows=zeilenzahl, cols=spalten)
    tabelle.style = "Table Grid"
    # Randlos: das Raster ordnet nur, es soll nicht zu sehen sein.
    tabelle.style = None
    tabelle.autofit = True

    for nummer, (lieferung, positionen) in enumerate(lieferungen):
        zelle = tabelle.cell(nummer % zeilenzahl, nummer // zeilenzahl)
        erster = True
        for text, farbe, groesse in _zeilen(lieferung, positionen, gross, normal):
            absatz = zelle.paragraphs[0] if erster else zelle.add_paragraph()
            erster = False
            _setze(absatz, text, farbe, groesse, eng=True)
        zelle.add_paragraph().paragraph_format.space_after = RASTER_ABSTAND

    return _als_bytes(doc)
