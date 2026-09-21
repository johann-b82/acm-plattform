"""Das ATR-Dokument aus dem Gerüst der Vorlage erzeugen.

**Die Vorlage ist der Rahmen.** Kopfblock, Tabellenüberschrift, Summenzeile
und Zertifizierungsblock bleiben, wie sie sind — samt Formaten, verbundenen
Zellen und Druckkopfzeile. Ersetzt wird nur der Bereich zwischen
Tabellenüberschrift und Summenzeile: dort stehen danach die Positionen der
Lieferung, nach Kategorie gruppiert. Eine Position ohne Katalogteil wird rot
hinterlegt, damit sie in Excel auffällt.

**Zeilen werden über Beschriftungen gefunden, nicht über Nummern.** An den
echten Vorlagen nachgesehen: die Tabellenüberschrift `PO Pos` steht in der
A350-Vorlage in Zeile 13, in der A380-Vorlage in Zeile 10; `Total weight`
in Zeile 80 bzw. 14. Feste Zeilennummern gingen bei der zweiten Vorlage
sofort daneben.

Aus `lumeapps` übernommen (`backend/app/services/atr_generate_xlsx.py`),
unverändert im Verhalten.
"""
from __future__ import annotations

import re
import zipfile
from copy import copy
from decimal import Decimal
from io import BytesIO

from openpyxl import load_workbook
from openpyxl.styles import Alignment, PatternFill
from openpyxl.worksheet.properties import PageSetupProperties

from app.atr.format import bestellposition, deutsches_datum, dokumentnummer, heute
from app.dokumente.blatt import auf_a4

ROT = PatternFill(start_color="FFFF0000", end_color="FFFF0000", fill_type="solid")
SPALTEN = 14  # A..N
#: Deutsches Dezimalkomma erzwingen, unabhängig davon, in welcher Sprache
#: LibreOffice beim Umwandeln läuft.
ZAHLFORMAT = "[$-407]0.00"
ACM_ANSCHRIFT = "ACM GmbH - Brandstücken 16 - 22549 Hamburg"
#: Der Zertifizierungstext bricht auf drei Zeilen um; verbundene Zellen wachsen
#: nicht von selbst mit.
ZERTIFIKAT_HOEHE = 48.0


class VorlageUnbrauchbar(ValueError):
    """Im Gerüst fehlt eine Zeile, ohne die nicht gesetzt werden kann."""


def _sichtbares_blatt(mappe):
    sichtbar = [b for b in mappe.worksheets if b.sheet_state == "visible"]
    if len(sichtbar) != 1:
        raise VorlageUnbrauchbar(
            f"Genau ein sichtbares Blatt erwartet, {len(sichtbar)} gefunden."
        )
    return sichtbar[0]


def _zeile_mit(blatt, spalte: int, text: str, bis: int = 20) -> int | None:
    """Erste Zeile, in der `spalte` mit `text` beginnt (ohne Rücksicht auf
    Groß- und Kleinschreibung)."""
    for zeile in range(1, min(bis, blatt.max_row) + 1):
        wert = blatt.cell(zeile, spalte).value
        if wert and str(wert).strip().lower().startswith(text.lower()):
            return zeile
    return None


def _summenzeile(blatt, ab: int) -> int:
    for zeile in range(ab, blatt.max_row + 1):
        wert = blatt.cell(zeile, 6).value
        if wert and "total" in str(wert).lower():
            return zeile
    raise VorlageUnbrauchbar('Im Gerüst fehlt die Zeile „Total weight“ in Spalte F.')


def _formate(blatt, zeile: int) -> list:
    return [copy(blatt.cell(zeile, s)._style) for s in range(1, SPALTEN + 1)]


def baue_atr(gerüst: bytes, lieferung: dict, positionen: list[dict]) -> bytes:
    """Setzt eine Lieferung in das Gerüst und gibt die fertige Mappe zurück."""
    mappe = load_workbook(BytesIO(gerüst))
    blatt = _sichtbares_blatt(mappe)

    ist_a380 = "380" in (lieferung.get("programm") or "")

    kopfzeile = _zeile_mit(blatt, 1, "PO Pos")
    if kopfzeile is None:
        raise VorlageUnbrauchbar('Im Gerüst fehlt die Tabellenüberschrift „PO Pos“.')
    erste_position = kopfzeile + 1

    # --- Kopfblock: nur die Felder der Lieferung, alles andere bleibt --------
    if (z := _zeile_mit(blatt, 1, "Supplier:")) is not None:
        blatt.cell(z, 4, ACM_ANSCHRIFT)
    if (z := _zeile_mit(blatt, 1, "Manufacturing Process Reference")) is not None:
        if lieferung.get("ba_auftrag"):
            blatt.cell(z, 4, lieferung["ba_auftrag"])
    if (z := _zeile_mit(blatt, 6, "Purchase Order No")) is not None:
        if lieferung.get("bestellnummer"):
            blatt.cell(z, 7, lieferung["bestellnummer"])
    if (z := _zeile_mit(blatt, 6, "MSN:")) is not None:
        if ist_a380:
            # A380 läuft als Ersatzteil, ohne MSN.
            blatt.cell(z, 7, lieferung.get("msn") or "N/A Spare Part")
        elif lieferung.get("msn"):
            blatt.cell(z, 7, lieferung["msn"])
            # Die A350-Bestellzeile steht dreigeteilt in I1/J1/K1
            # („ 79 - " | MSN | "- 94"); die MSN dort mit Nullen aufgefüllt.
            blatt.cell(1, 10, str(lieferung["msn"]).zfill(4))
    wiegezeile = _zeile_mit(blatt, 1, "Weighing date")
    if wiegezeile is not None:
        # Auch ohne gesetztes Datum wird geschrieben: die Vorlage trägt dort
        # `=TODAY()`, und LibreOffice setzt das in seiner eigenen Sprache —
        # im PDF stand „9/10/2026" statt „10.09.2026". Nachgemessen an einer
        # erzeugten Datei. Ein Dokument, das heute entsteht, darf heute
        # tragen; es muss nur deutsch lesbar sein.
        blatt.cell(
            wiegezeile, 3, deutsches_datum(lieferung.get("wiegedatum") or heute())
        )
        blatt.cell(
            wiegezeile, 12, deutsches_datum(lieferung.get("pruefdatum") or heute())
        )
        # Die Satzbezeichnung steht als Banner direkt darüber.
        if lieferung.get("satz_titel"):
            blatt.cell(wiegezeile - 1, 1, lieferung["satz_titel"])

    # Druckkopfzeile rechts. Der PDF-Weg baut sie noch einmal auf; hier steht
    # sie, damit ein direkt aus Excel gedrucktes Blatt dasselbe zeigt.
    blatt.oddHeader.right.text = (
        f"Doc-No.: {dokumentnummer(lieferung.get('atr_nummer'), lieferung.get('programm'))}\n"
        f"Date: {heute().strftime('%d.%m.%Y')}\n"
        "Page: &P of &N"
    )

    # --- Formate merken, bevor der Bereich umgebaut wird ---------------------
    positionsformat = _formate(blatt, erste_position + 1)
    abschnittsformat = _formate(blatt, erste_position)

    summe_alt = _summenzeile(blatt, erste_position)
    alter_bereich = max(0, summe_alt - erste_position)

    # `delete_rows`/`insert_rows` verschieben Zellinhalte, aber nicht die
    # verbundenen Bereiche — die blieben liegen, wo sie waren, und der breite
    # Zertifizierungsblock rutschte auseinander. Deshalb vorher merken.
    alte_verbuende = [
        (b.min_row, b.min_col, b.max_row, b.max_col)
        for b in blatt.merged_cells.ranges
    ]

    # Nach Kategorie gruppieren, in der Reihenfolge des ersten Auftretens.
    gruppen: list[tuple[str | None, list[dict]]] = []
    wo: dict[str | None, int] = {}
    for p in positionen:
        kategorie = p.get("kategorie")
        if kategorie not in wo:
            wo[kategorie] = len(gruppen)
            gruppen.append((kategorie, []))
        gruppen[wo[kategorie]][1].append(p)

    neue_zeilen = sum(1 + len(liste) for _, liste in gruppen)

    if alter_bereich:
        blatt.delete_rows(erste_position, alter_bereich)
    if neue_zeilen:
        blatt.insert_rows(erste_position, neue_zeilen)

    # Verbünde nachziehen: was über dem Positionsbereich lag, bleibt; was darin
    # lag, fällt weg; was ab der Summenzeile kam, rückt um die Differenz.
    versatz = neue_zeilen - alter_bereich
    blatt.merged_cells.ranges = []
    for z1, s1, z2, s2 in alte_verbuende:
        if z1 < erste_position:
            blatt.merge_cells(start_row=z1, start_column=s1, end_row=z2, end_column=s2)
        elif z1 >= summe_alt:
            blatt.merge_cells(
                start_row=z1 + versatz, start_column=s1,
                end_row=z2 + versatz, end_column=s2,
            )

    zeile = erste_position
    gesamt = Decimal("0")
    for kategorie, liste in gruppen:
        for spalte in range(1, SPALTEN + 1):
            blatt.cell(zeile, spalte)._style = copy(abschnittsformat[spalte - 1])
        blatt.cell(zeile, 1, kategorie or "")
        # Die Abschnittsüberschrift läuft über A..H, die Prüfspalten I..N
        # bleiben getrennt.
        blatt.merge_cells(start_row=zeile, start_column=1, end_row=zeile, end_column=8)
        blatt.cell(zeile, 1).alignment = Alignment(horizontal="center", vertical="center")
        zeile += 1

        for p in liste:
            for spalte in range(1, SPALTEN + 1):
                blatt.cell(zeile, spalte)._style = copy(positionsformat[spalte - 1])
            blatt.cell(zeile, 1, bestellposition(p.get("bestellposition")) or "")
            blatt.cell(zeile, 2, p.get("lieferantennummer") or "")
            blatt.cell(zeile, 3, p.get("teilenummer") or "")
            blatt.cell(zeile, 4, p.get("bezeichnung") or "")
            serien = p.get("seriennummern") or []
            blatt.cell(zeile, 5, ", ".join(serien) if serien else "N/A")
            blatt.cell(zeile, 6, p.get("zeichnung") or "")
            menge = p.get("menge") or 1
            blatt.cell(zeile, 7, menge)
            if p.get("gewicht_kg") is not None:
                # Das gespeicherte Gewicht gilt je Stück; Zeile und Summe
                # rechnen mit der Menge.
                zeilengewicht = Decimal(str(p["gewicht_kg"])) * menge
                zelle = blatt.cell(zeile, 8, float(zeilengewicht))
                zelle.number_format = ZAHLFORMAT
                gesamt += zeilengewicht
            for spalte in range(9, 14):
                blatt.cell(zeile, spalte, "P")
            blatt.cell(zeile, 14, "OK")
            if not p.get("teil_id"):
                for spalte in range(1, SPALTEN + 1):
                    blatt.cell(zeile, spalte).fill = ROT
            zeile += 1

    summe_neu = _summenzeile(blatt, erste_position)
    zelle = blatt.cell(summe_neu, 8, float(gesamt))
    zelle.number_format = ZAHLFORMAT

    # Das Hoechstgewicht ist meist eine Konstante des Programms und steht
    # schon in der Vorlage; ueberschrieben wird nur, wenn die Lieferung einen
    # eigenen Wert traegt. Das **Format** wird in jedem Fall angeglichen: sonst
    # steht auf demselben Blatt „211.00" neben „4,63".
    for z in range(summe_neu, min(summe_neu + 4, blatt.max_row + 1)):
        if "max" in str(blatt.cell(z, 6).value or "").lower():
            zelle = blatt.cell(z, 8)
            if lieferung.get("max_gewicht_kg") is not None:
                zelle.value = float(lieferung["max_gewicht_kg"])
            if isinstance(zelle.value, (int, float)):
                zelle.number_format = ZAHLFORMAT

    # Die Zertifizierungszeile bricht auf drei Zeilen um; verbundene Zellen
    # wachsen nicht mit. Erst nach dem Umbau suchen — der Block ist gewandert.
    zertifikat = next(
        (
            z.row
            for reihe in blatt.iter_rows()
            for z in reihe
            if isinstance(z.value, str) and "hereby certify" in z.value
        ),
        None,
    )
    if zertifikat is not None:
        blatt.row_dimensions[zertifikat].height = ZERTIFIKAT_HOEHE

    # QS-Unterschrift auf der „Date:"-Zeile. Die Vorlage bringt einen festen
    # Namen mit; der Wert der Lieferung ist maßgeblich, auch wenn er leer ist —
    # sonst stünde dort ein Name, den niemand gesetzt hat.
    unterschrift = next(
        (
            z
            for z in range(summe_neu, blatt.max_row + 1)
            if str(blatt.cell(z, 1).value or "").strip() == "Date:"
        ),
        None,
    )
    if unterschrift is not None:
        blatt.cell(unterschrift, 3, lieferung.get("qs_unterschrift") or "")
        blatt.cell(unterschrift, 12, lieferung.get("qs_unterschrift") or "")
        # Die Vorlage trägt dort `=TODAY()`, was LibreOffice in seiner eigenen
        # Sprache setzt (8/5/2026). Als feste Zeichenkette sind alle vier
        # Datumsfelder des Blatts gleich formatiert.
        datum = deutsches_datum(
            lieferung.get("pruefdatum") or lieferung.get("wiegedatum") or heute()
        )
        blatt.cell(unterschrift, 2, datum)
        blatt.cell(unterschrift, 8, datum)

    # Auf eine Seite **Breite** zwingen; die Höhe darf laufen. Ohne das
    # verdoppelte der Überlauf nach rechts die Seitenzahl.
    blatt.print_area = f"A1:N{blatt.max_row}"
    auf_a4(blatt, quer=True)
    blatt.page_setup.fitToWidth = 1
    blatt.page_setup.fitToHeight = 0
    # Gleiche Ränder links und rechts, damit die rechtsbündige Druckkopfzeile
    # über der rechten Kante der Tabelle steht. Die Vorlage bringt ungleiche
    # Ränder mit und schiebt die Kopfzeile darüber hinaus.
    blatt.page_margins.right = blatt.page_margins.left
    blatt.page_setup.scale = None  # schließt sich mit fitToPage aus
    blatt.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)

    puffer = BytesIO()
    mappe.save(puffer)
    return _nachbehandeln(puffer.getvalue(), gerüst)


# --- Nachbehandlung der gespeicherten Mappe -------------------------------------------------
# openpyxl behält den Kopfzeilen-Code `&G`, wirft die zugehörigen Bildteile beim
# Speichern aber weg — ein gedrucktes Blatt zeigt dann kein Logo.
#
# Die Teile kommen aus der Vorlage zurück, und die Dateien des Ergebnisses
# werden **ergänzt**, nicht durch die der Vorlage ersetzt. Der Unterschied ist
# nicht akademisch: openpyxl schreibt eine andere Menge von Teilen als Excel
# (etwa ohne `sharedStrings.xml`). Ein `[Content_Types].xml` aus der Vorlage
# verspräche dann Teile, die es nicht gibt, und die Datei ließe sich nicht mehr
# öffnen.
BEZIEHUNG_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PAKET_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
KENNUNG = "rIdHF1"


def _nachbehandeln(fertig: bytes, gerüst: bytes) -> bytes:
    """Zwei Dinge, die openpyxl beim Speichern verliert oder verdreht.

    Die Umbrüche der Druckkopfzeile werden **immer** geradegezogen; das
    Kopfbild nur, wenn die Vorlage eines hat. Beides hing hier einmal am
    selben Zweig — ein Gerüst ohne Logo bekam dann weiterhin die kaputte
    Kopfzeile.
    """
    try:
        with zipfile.ZipFile(BytesIO(gerüst)) as vorlage:
            bild = vorlage.read("xl/media/image1.png")
            vml = vorlage.read("xl/drawings/vmlDrawing1.vml")
            vml_rels = vorlage.read("xl/drawings/_rels/vmlDrawing1.vml.rels")
    except KeyError:
        bild = vml = vml_rels = None

    with zipfile.ZipFile(BytesIO(fertig)) as neu:
        namen = neu.namelist()
        # Das Blatt, dessen Kopfzeile auf ein Bild verweist — nur dort gehören
        # die Bildteile hin.
        mit_bild = (
            next(
                (
                    n
                    for n in namen
                    if n.startswith("xl/worksheets/sheet")
                    and n.endswith(".xml")
                    and b"&amp;G" in neu.read(n)
                ),
                None,
            )
            if bild is not None
            else None
        )
        rels_name = (
            f"xl/worksheets/_rels/{mit_bild.split('/')[-1]}.rels" if mit_bild else None
        )

        ergebnis = BytesIO()
        with zipfile.ZipFile(ergebnis, "w", zipfile.ZIP_DEFLATED) as aus:
            for name in namen:
                daten = neu.read(name)
                if name.startswith("xl/worksheets/sheet") and name.endswith(".xml"):
                    daten = _kopfzeile_umbrueche(daten.decode("utf-8")).encode("utf-8")
                    if name == mit_bild:
                        daten = _mit_kopfzeichnung(daten)
                elif mit_bild and name == "[Content_Types].xml":
                    daten = _mit_typen(daten)
                elif mit_bild and name == rels_name:
                    daten = _mit_vml_beziehung(daten)
                aus.writestr(name, daten)

            if mit_bild:
                aus.writestr("xl/media/image1.png", bild)
                aus.writestr("xl/drawings/vmlDrawing1.vml", vml)
                aus.writestr("xl/drawings/_rels/vmlDrawing1.vml.rels", vml_rels)
                if rels_name not in namen:
                    aus.writestr(
                        rels_name,
                        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                        f'<Relationships xmlns="{PAKET_NS}">'
                        f'<Relationship Id="{KENNUNG}" Type="{BEZIEHUNG_NS}/vmlDrawing" '
                        'Target="../drawings/vmlDrawing1.vml"/></Relationships>',
                    )
    return ergebnis.getvalue()


def _mit_typen(content_types: bytes) -> bytes:
    """`png` und `vml` als bekannte Endungen ergänzen, falls sie fehlen."""
    text = content_types.decode("utf-8")
    zusatz = ""
    if 'Extension="png"' not in text:
        zusatz += '<Default Extension="png" ContentType="image/png"/>'
    if 'Extension="vml"' not in text:
        zusatz += (
            '<Default Extension="vml" ContentType='
            '"application/vnd.openxmlformats-officedocument.vmlDrawing"/>'
        )
    if not zusatz:
        return content_types
    return re.sub(r"(<Types[^>]*>)", r"\1" + zusatz, text, count=1).encode("utf-8")


def _mit_kopfzeichnung(blatt: bytes) -> bytes:
    text = blatt.decode("utf-8")
    if "legacyDrawingHF" in text:
        return text.encode("utf-8")
    if "xmlns:r=" not in text.split(">", 1)[0]:
        text = text.replace("<worksheet ", f'<worksheet xmlns:r="{BEZIEHUNG_NS}" ', 1)
    return text.replace(
        "</worksheet>", f'<legacyDrawingHF r:id="{KENNUNG}"/></worksheet>', 1
    ).encode("utf-8")


def _kopfzeile_umbrueche(blatt_xml: str) -> str:
    """Zeilenumbrüche in der Druckkopfzeile echt machen.

    openpyxl schreibt ein `\n` in der Kopfzeile als `_x000a_`. Excel versteht
    das noch, LibreOffice nicht — dort steht dann wörtlich
    „…Issue: 01_x000a_Date: 10.09.2026" quer über der Seite. Nachgemessen an
    einer erzeugten Mappe, bevor diese Zeile hier stand.
    """
    if "_x000a_" not in blatt_xml:
        return blatt_xml
    anfang = blatt_xml.find("<headerFooter")
    if anfang == -1:
        return blatt_xml
    ende = blatt_xml.find("</headerFooter>", anfang)
    if ende == -1:
        return blatt_xml
    ende += len("</headerFooter>")
    return (
        blatt_xml[:anfang]
        + blatt_xml[anfang:ende].replace("_x000a_", "&#10;")
        + blatt_xml[ende:]
    )


def _mit_vml_beziehung(rels: bytes) -> bytes:
    text = rels.decode("utf-8")
    if "vmlDrawing" in text:
        return rels
    eintrag = (
        f'<Relationship Id="{KENNUNG}" Type="{BEZIEHUNG_NS}/vmlDrawing" '
        'Target="../drawings/vmlDrawing1.vml"/>'
    )
    return text.replace("</Relationships>", eintrag + "</Relationships>", 1).encode("utf-8")
