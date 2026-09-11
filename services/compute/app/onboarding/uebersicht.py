"""Die Schulungsübersicht als Formblatt 71.

Das Papierformular, das ein neuer Mitarbeiter bekommt: je Schulung eine Zeile
mit laufender Nummer, Zeitraum, Bezeichnung und zwei Kreuzfeldpaaren — intern
oder extern, Nachweis vorhanden ja oder nein.

**Für einen Neueintritt bleiben Zeitraum und Kreuze leer.** Das Blatt ist dann
der Plan, den er abarbeitet. Ein Datum vorzugeben, das noch niemand terminiert
hat, wäre erfunden — und stünde hinterher gedruckt im Ordner.

Wie beim Einarbeitungsbogen: openpyxl schreibt eine Mappe, LibreOffice macht
ein PDF daraus. Alles steht in Zellen, kein UNO nötig.
"""
from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO

from openpyxl import Workbook
from openpyxl.drawing.image import Image as ExcelBild
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.worksheet.page import PageMargins

from app.dokumente import qr as qr_mod
from app.dokumente.blatt import auf_a4
from app.dokumente.logo import Logo
from app.dokumente.pdf import nach_pdf

#: Kopfangaben des Formblatts. Stehen so auf dem Papier und ändern sich nur
#: mit einer neuen Revision.
FORMBLATT = "Formblatt 71"
REVISIONS_INDEX = "A"
REVISIONS_STAND = "22.03.2022"
AUSGABEDATUM = "22.03.2022"

TITEL = "Schulungsübersicht"

#: Die vier Ankreuzspalten (D–G). Vier eigene Felder, nicht zwei kombinierte:
#: das Blatt wird gedruckt und von Hand abgehakt.
SP_IN, SP_EX, SP_JA, SP_NEIN = 4, 5, 6, 7
SPALTENBREITEN = (11, 16, 58, 5, 5, 5, 6)

#: Wo QR und Passermarken sitzen, wenn das Blatt Teil eines Vorgangs ist.
#: Spalte A ist breit genug für eine Marke, Spalte G trägt den QR.
QR_SPALTE = 7
GEOMETRIE = qr_mod.Geometrie(
    spaltenbreiten=SPALTENBREITEN, rand_links_px=48.0, rand_oben_pt=43.2
)

_DUENN = Side(style="thin", color="000000")
_RAHMEN = Border(left=_DUENN, right=_DUENN, top=_DUENN, bottom=_DUENN)
_OBEN_LINKS = Alignment(horizontal="left", vertical="top", wrap_text=True)
_MITTE = Alignment(horizontal="center", vertical="center")
_MITTE_UMBRUCH = Alignment(horizontal="center", vertical="center", wrap_text=True)
_MITTE_OBEN = Alignment(horizontal="center", vertical="top")


@dataclass
class Zeile:
    """Eine Zeile des Formblatts."""

    bezeichnung: str
    #: Freitext wie „16.03.2026 – 17.03.2026"; leer, solange nicht absolviert.
    zeitraum: str = ""
    #: Anbieter samt Anschrift; steht unter der Bezeichnung, nicht in einer
    #: eigenen Spalte — so ist das Formular gebaut.
    anbieter: str = ""
    #: True = intern (IN), False = extern (EX), None = noch offen.
    intern: bool | None = None
    #: Setzt das Kreuz bei „Schulungsnachweis vorhanden". None = offen.
    nachweis: bool | None = None


def _kopf(blatt, name: str, funktion: str, logo: Logo | None) -> int:
    """Titelblock und Personenangaben. Gibt die nächste freie Zeile zurück."""
    # „Blatt x von y" als Seitenkopf: fest verdrahtet wäre es falsch, sobald
    # die Liste auf eine zweite Seite läuft.
    blatt.oddHeader.right.text = "&9Blatt &P von &N"
    blatt.evenHeader.right.text = blatt.oddHeader.right.text

    oben = 1
    if logo is not None:
        breite = 110
        bild = ExcelBild(BytesIO(logo.daten))
        bild.width = breite
        bild.height = max(1, round(breite * logo.hoehe / logo.breite))
        blatt.add_image(bild, "A1")
        for zeile in (1, 2, 3):
            blatt.row_dimensions[zeile].height = 20
        oben = 4

    for i, text in enumerate(
        (FORMBLATT, f"Revisions-Index: {REVISIONS_INDEX}", f"Revisions-Stand: {REVISIONS_STAND}")
    ):
        zeile = oben + i
        blatt.merge_cells(start_row=zeile, start_column=3, end_row=zeile, end_column=SP_NEIN)
        zelle = blatt.cell(zeile, 3, text)
        zelle.font = Font(size=9)
        zelle.alignment = Alignment(horizontal="right", vertical="center")

    titel = blatt.cell(oben, 1, TITEL)
    titel.font = Font(size=16, bold=True)
    titel.alignment = Alignment(horizontal="left", vertical="center")
    blatt.merge_cells(start_row=oben, start_column=1, end_row=oben + 1, end_column=2)

    naechste = oben + 4
    blatt.cell(naechste, 1, "Name:").font = Font(bold=True)
    blatt.cell(naechste, 2, name)
    blatt.cell(naechste + 1, 1, "Funktion:").font = Font(bold=True)
    blatt.cell(naechste + 1, 2, funktion or "—")
    for zeile in (naechste, naechste + 1):
        blatt.merge_cells(start_row=zeile, start_column=2, end_row=zeile, end_column=SP_NEIN)
    return naechste + 2


def _tabellenkopf(blatt, zeile: int) -> int:
    """Zweizeiliger Tabellenkopf wie auf dem Formular."""
    blatt.cell(zeile, 1, "Laufende\nNummer:")
    blatt.cell(zeile, 2, "durchgeführt am/\nvon … bis:")
    blatt.cell(zeile, 3, "Bezeichnung der Aus- und Fortbildungsmaßnahme")
    blatt.cell(zeile, SP_IN, "Schulungsnachweis\nvorhanden")
    blatt.merge_cells(start_row=zeile, start_column=SP_IN, end_row=zeile, end_column=SP_NEIN)

    unten = zeile + 1
    blatt.cell(unten, 3, "Interne (IN) oder externe (EX) Maßnahme:")
    blatt.cell(unten, SP_IN, "IN")
    blatt.cell(unten, SP_EX, "EX")
    blatt.cell(unten, SP_JA, "ja")
    blatt.cell(unten, SP_NEIN, "nein")
    for spalte in (1, 2):
        blatt.merge_cells(start_row=zeile, start_column=spalte, end_row=unten, end_column=spalte)

    for r in (zeile, unten):
        for c in range(1, SP_NEIN + 1):
            zelle = blatt.cell(r, c)
            zelle.font = Font(size=9, bold=True)
            # Umbruch auch in den zentrierten Spalten — sonst läuft
            # „Schulungsnachweis vorhanden" in einer Zeile durch.
            zelle.alignment = _MITTE_UMBRUCH if c >= SP_IN else _OBEN_LINKS
            zelle.border = _RAHMEN
    blatt.row_dimensions[zeile].height = 32
    return unten + 1


def _zeile_schreiben(blatt, r: int, nummer: int, z: Zeile) -> None:
    blatt.cell(r, 1, f"{nummer:02d}")
    blatt.cell(r, 2, z.zeitraum)
    blatt.cell(r, 3, z.bezeichnung + (f"\n\n{z.anbieter}" if z.anbieter else ""))
    # Kreuz nur, wo die Angabe feststeht; offen bleibt leer zum Ankreuzen.
    blatt.cell(r, SP_IN, "X" if z.intern is True else "")
    blatt.cell(r, SP_EX, "X" if z.intern is False else "")
    blatt.cell(r, SP_JA, "X" if z.nachweis is True else "")
    blatt.cell(r, SP_NEIN, "X" if z.nachweis is False else "")

    for c in range(1, SP_NEIN + 1):
        zelle = blatt.cell(r, c)
        zelle.border = _RAHMEN
        zelle.font = Font(size=9)
        zelle.alignment = _MITTE_OBEN if c == 1 else (_MITTE if c >= SP_IN else _OBEN_LINKS)
    blatt.row_dimensions[r].height = 38 if z.anbieter else 24


def _fuss(blatt, freigegeben_von: str, erstellt_von: str) -> None:
    """Der Fußblock als echte Seitenfußzeile.

    Nicht als Tabellenzeilen unter der letzten Schulung: dort klebte er am
    Tabellenende statt am Blattfuß und stünde bei wenigen Zeilen mitten auf der
    Seite.

    Bewusst je Abschnitt **eine** Zeile: openpyxl kodiert ein „\\n" in der
    Fußzeile als OOXML-Escape `_x000a_`, und LibreOffice gibt das wörtlich aus
    statt umzubrechen.
    """
    blatt.oddFooter.left.text = (
        f"&9Aktualisiert am: ________     Freigegeben von: {freigegeben_von or '________'}"
    )
    blatt.oddFooter.right.text = (
        f"&9Ausgabedatum: {AUSGABEDATUM}     Erstellt von: {erstellt_von or '________'}"
    )
    blatt.evenFooter.left.text = blatt.oddFooter.left.text
    blatt.evenFooter.right.text = blatt.oddFooter.right.text


def fuelle_blatt(
    blatt,
    name: str,
    funktion: str,
    zeilen: list[Zeile],
    *,
    freigegeben_von: str = "",
    erstellt_von: str = "",
    logo: Logo | None = None,
    doc_uid: str | None = None,
    layout_raus: dict | None = None,
) -> None:
    """Formblatt 71 in ein vorhandenes Arbeitsblatt schreiben."""
    blatt.title = "Schulungsübersicht"
    for spalte, breite in zip("ABCDEFG", SPALTENBREITEN):
        blatt.column_dimensions[spalte].width = breite

    felder: list | None = [] if layout_raus is not None else None

    r = _kopf(blatt, name, funktion, logo)
    kopfzeile = r
    qr_zeile = kopfzeile
    r = _tabellenkopf(blatt, r)
    for i, z in enumerate(zeilen, start=1):
        _zeile_schreiben(blatt, r, i, z)
        if felder is not None:
            # Ein Pflichtfeld je Zeile: der Zeitraum. Er sagt, dass die
            # Schulung stattgefunden hat; die Kreuze daneben sagen nur, wie.
            felder.append((f"zeitraum_{i}", f"Zeitraum — {z.bezeichnung[:40]}", r, 2, 2))
        r += 1
    letzte = r - 1
    _fuss(blatt, freigegeben_von, erstellt_von)

    qr_mod.hoehen_festschreiben(blatt, max(letzte, kopfzeile + 2))

    if doc_uid:
        blatt.row_dimensions[qr_zeile].height = 36
        qr_mod.qr_einsetzen(blatt, doc_uid, f"G{qr_zeile}")
        qr_mod.marke_einsetzen(blatt, "A1")
        qr_mod.marke_einsetzen(blatt, f"A{max(letzte, kopfzeile + 2)}")

    blatt.print_area = f"A1:G{max(r - 1, kopfzeile)}"
    auf_a4(blatt)
    blatt.page_setup.fitToWidth = 1
    blatt.page_setup.fitToHeight = 0
    blatt.sheet_properties.pageSetUpPr.fitToPage = True
    # Unten mehr Rand, damit die Fußzeile nicht an der Tabelle klebt.
    blatt.page_margins = PageMargins(
        left=0.5, right=0.4, top=0.6, bottom=0.8, header=0.3, footer=0.35
    )
    # Läuft die Liste auf eine zweite Seite, wiederholt sich der Tabellenkopf.
    blatt.print_title_rows = f"{kopfzeile}:{kopfzeile + 1}"

    if layout_raus is not None:
        layout_raus.update(
            qr_mod.layout(
                blatt,
                GEOMETRIE,
                doc_uid=doc_uid or "",
                felder=felder or [],
                qr_spalte=QR_SPALTE,
                qr_zeile=qr_zeile,
                marken=[(1, 1), (1, max(letzte, kopfzeile + 2))],
            )
        )


def baue_xlsx(
    name: str,
    funktion: str,
    zeilen: list[Zeile],
    *,
    freigegeben_von: str = "",
    erstellt_von: str = "",
    logo: Logo | None = None,
    doc_uid: str | None = None,
    layout_raus: dict | None = None,
) -> bytes:
    mappe = Workbook()
    fuelle_blatt(
        mappe.active,
        name,
        funktion,
        zeilen,
        freigegeben_von=freigegeben_von,
        erstellt_von=erstellt_von,
        logo=logo,
        doc_uid=doc_uid,
        layout_raus=layout_raus,
    )
    puffer = BytesIO()
    mappe.save(puffer)
    return puffer.getvalue()


async def baue_pdf(
    name: str,
    funktion: str,
    zeilen: list[Zeile],
    *,
    freigegeben_von: str = "",
    erstellt_von: str = "",
    logo: Logo | None = None,
    doc_uid: str | None = None,
    layout_raus: dict | None = None,
) -> bytes:
    return await nach_pdf(
        baue_xlsx(
            name,
            funktion,
            zeilen,
            freigegeben_von=freigegeben_von,
            erstellt_von=erstellt_von,
            logo=logo,
            doc_uid=doc_uid,
            layout_raus=layout_raus,
        ),
        name="schulungsuebersicht",
    )
