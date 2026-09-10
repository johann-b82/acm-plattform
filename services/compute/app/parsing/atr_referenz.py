"""Parser für die ATR-Referenzmappe (Teilekatalog + Kopfdaten der Vorlage).

Die Mappe ist ein ausgefülltes ATR-Formular: Kopfdaten stehen in festen
Zellen, ab Zeile 14 folgen die Teile, unterbrochen von Abschnittsüberschriften
in Spalte A. Ein Teil erkennt man daran, dass in Spalte C eine Nummer mit
Präfix ``VR`` steht.

Diese Zelladressen sind aus dem Altprojekt übernommen
(`backend/app/services/atr_reference_import.py`) und an derselben Vorlage
gewachsen. Sie sind der Grund, warum vor dem ersten Import ein echtes
Referenzblatt durchlaufen sollte: verschiebt der Kunde eine Zeile, findet der
Parser die Kopfdaten am falschen Ort — und die Prüfung auf Zeile 13 ist genau
dafür da, das laut scheitern zu lassen statt still falsche Werte zu übernehmen.
"""
from __future__ import annotations

from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation
from io import BytesIO

from openpyxl import load_workbook

#: Kopfdaten: Feldname → Zelle.
KOPFZELLEN = {
    "kunde": "D1",
    "programm": "D2",
    "arbeitspaket": "D3",
    "besteller_spez": "D4",
    "atp": "D5",
    "lieferanten_spez": "D6",
    "referenz": "D7",
    "lieferant": "D8",
    "kunden_spez": "G8",
    "nscm": "G4",
    "ata_kapitel": "G3",
    "waage": "F12",
}


@dataclass
class Teil:
    reihenfolge: int
    teilenummer: str
    lieferantennummer: str | None
    bezeichnung: str | None
    zeichnung: str | None
    menge: int
    gewicht_kg: Decimal | None
    kategorie: str | None


@dataclass
class Referenzmappe:
    kopf: dict[str, str | None]
    teile: list[Teil] = field(default_factory=list)
    hinweise: list[str] = field(default_factory=list)


class MappeUnbrauchbar(ValueError):
    """Die Datei sieht nicht aus wie ein ATR-Referenzblatt."""


def _zelle(blatt, adresse: str) -> str | None:
    wert = blatt[adresse].value
    if wert is None:
        return None
    text = str(wert).strip()
    return text or None


def _dezimal(wert) -> Decimal | None:
    if wert is None:
        return None
    text = str(wert).strip().replace(",", ".")
    if not text:
        return None
    try:
        return Decimal(text)
    except (InvalidOperation, ValueError):
        return None


def _ganzzahl(wert, ersatz: int = 1) -> int:
    if wert is None:
        return ersatz
    try:
        return int(float(str(wert).strip().replace(",", ".")))
    except (TypeError, ValueError):
        return ersatz


def lies_referenzmappe(daten: bytes) -> Referenzmappe:
    mappe = load_workbook(BytesIO(daten), data_only=True, read_only=False)
    sichtbar = [b for b in mappe.worksheets if b.sheet_state == "visible"]
    if len(sichtbar) != 1:
        raise MappeUnbrauchbar(
            f"Genau ein sichtbares Blatt erwartet, {len(sichtbar)} gefunden."
        )
    blatt = sichtbar[0]

    # Der Riegel: passt die Kopfzeile der Teiletabelle nicht, stimmt das
    # Layout nicht, und alle Zelladressen oben zeigen ins Leere.
    c13 = (_zelle(blatt, "C13") or "").lower()
    h13 = (_zelle(blatt, "H13") or "").lower()
    if "part number" not in c13 or "weight" not in h13:
        raise MappeUnbrauchbar(
            "Unbekanntes ATR-Layout: Zeile 13 trägt nicht die erwartete "
            "Tabellenüberschrift."
        )

    kopf = {feld: _zelle(blatt, adresse) for feld, adresse in KOPFZELLEN.items()}

    teile: list[Teil] = []
    hinweise: list[str] = []
    kategorie: str | None = None
    reihenfolge = 0

    for zeile in range(14, blatt.max_row + 1):
        a = _zelle(blatt, f"A{zeile}")
        c = _zelle(blatt, f"C{zeile}")
        f = _zelle(blatt, f"F{zeile}")

        if f and "total" in f.lower():
            break  # Summenblock — danach kommen keine Teile mehr.

        if c and c.upper().startswith("VR"):
            roh = blatt[f"H{zeile}"].value
            gewicht = _dezimal(roh)
            if roh not in (None, "") and gewicht is None:
                hinweise.append(f"Zeile {zeile}: Gewicht {roh!r} nicht lesbar")
            reihenfolge += 1
            teile.append(
                Teil(
                    reihenfolge=reihenfolge,
                    teilenummer=c,
                    lieferantennummer=_zelle(blatt, f"B{zeile}"),
                    bezeichnung=_zelle(blatt, f"D{zeile}"),
                    zeichnung=f,
                    menge=_ganzzahl(blatt[f"G{zeile}"].value),
                    gewicht_kg=gewicht,
                    kategorie=kategorie,
                )
            )
        elif a and not c:
            # Abschnittsüberschrift. „If applicable" ist keine.
            if not a.lower().startswith("if applicable"):
                kategorie = a

    if not teile:
        hinweise.append("Keine Teilezeilen gefunden.")
    return Referenzmappe(kopf=kopf, teile=teile, hinweise=hinweise)
