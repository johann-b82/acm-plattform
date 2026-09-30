"""Die Interviewdatei lesen — ein Blatt je Bereich, Stufe 0–3.

Anders als die alte Qualifikationsmatrix ist die Interviewdatei nicht
transponiert: Zeilen sind Personen, Spalten sind Aufgabenfamilien. Der Aufbau je
Blatt (Kopf auf Zeile 4):

    Mitarbeiter | Personio Team | Produkte / Produktfamilien
                | <Aufgabenfamilie 1 … N> | Weitere Bereiche
                | Engpass / Trainingsbedarf | Validiert durch | Datum | Notiz

**Die Aufgabenfamilien sind genau die Spalten zwischen „Produkte /
Produktfamilien" und „Weitere Bereiche".** Damit fällt der `#REF!`-Fehler der
Datei (eine beschriebene, aber nicht vorhandene Spalte „Sonderarbeiten") von
selbst weg — was keine Spalte hat, ist keine Familie.

Unter dem Personenblock steht ein zweiter Block „Aufgabenfamilie | Enthält aus
bestehender Matrix" mit den Beschreibungen. Blätter ohne „Mitarbeiter"-Kopf
(Start, Capacity View, Tabelle1/2) werden übersprungen.
"""
from __future__ import annotations

from dataclasses import dataclass, field
from datetime import date, datetime
from io import BytesIO

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel


class InterviewUnbrauchbar(Exception):
    """Die Datei enthält kein einziges auswertbares Bereichsblatt."""


@dataclass
class Person:
    name: str
    team: str | None
    produkte: str | None
    weitere_bereiche: str | None
    engpass: str | None
    validiert_durch: str | None
    validiert_am: date | None
    notiz: str | None
    #: Aufgabenfamilie → Stufe 0–3 (nur gesetzte).
    stufen: dict[str, int] = field(default_factory=dict)


@dataclass
class Familie:
    name: str
    beschreibung: str | None


@dataclass
class Blatt:
    bereich: str
    familien: list[Familie] = field(default_factory=list)
    personen: list[Person] = field(default_factory=list)


@dataclass
class Datei:
    dateiname: str
    blaetter: list[Blatt] = field(default_factory=list)
    hinweise: list[str] = field(default_factory=list)


def _text(wert) -> str | None:
    if wert is None:
        return None
    s = str(wert).strip()
    return s or None


def _stufe(wert) -> int | None:
    if isinstance(wert, bool):
        return None
    if isinstance(wert, (int, float)):
        z = int(wert)
        return z if 0 <= z <= 3 else None
    return None


def _datum(wert) -> date | None:
    if isinstance(wert, datetime):
        return wert.date()
    if isinstance(wert, date):
        return wert
    if isinstance(wert, (int, float)):
        try:
            gewandelt = from_excel(wert)
        except (ValueError, OverflowError):
            return None
        return gewandelt.date() if isinstance(gewandelt, datetime) else gewandelt
    return None


def _kopfzeile(ws) -> int | None:
    """Die Zeile mit „Mitarbeiter" in Spalte A (Zeile 1–8), oder nichts."""
    for r in range(1, 9):
        if _text(ws.cell(r, 1).value) == "Mitarbeiter":
            return r
    return None


def _spalten(ws, kopf: int) -> dict[str, int]:
    """Beschriftung → Spaltennummer der Kopfzeile."""
    spalten: dict[str, int] = {}
    for c in range(1, ws.max_column + 1):
        text = _text(ws.cell(kopf, c).value)
        if text:
            spalten.setdefault(text, c)
    return spalten


def _finde(spalten: dict[str, int], teil: str) -> int | None:
    for text, c in spalten.items():
        if teil.lower() in text.lower():
            return c
    return None


def _lies_blatt(ws) -> Blatt | None:
    kopf = _kopfzeile(ws)
    if kopf is None:
        return None
    spalten = _spalten(ws, kopf)
    c_produkte = _finde(spalten, "Produkte")
    c_weitere = _finde(spalten, "Weitere Bereiche")
    if c_produkte is None or c_weitere is None or c_weitere <= c_produkte + 1:
        return None

    # Aufgabenfamilien: die Spalten strikt zwischen Produkte und Weitere Bereiche.
    familien_spalten: list[tuple[str, int]] = []
    for c in range(c_produkte + 1, c_weitere):
        text = _text(ws.cell(kopf, c).value)
        if text:
            familien_spalten.append((text, c))
    if not familien_spalten:
        return None

    c_team = _finde(spalten, "Personio Team")
    c_engpass = _finde(spalten, "Engpass")
    c_validiert = _finde(spalten, "Validiert")
    c_datum = spalten.get("Datum")
    c_notiz = spalten.get("Notiz")

    blatt = Blatt(bereich=ws.title.strip())

    # Personen: ab Kopf+1, bis Spalte A leer ist.
    r = kopf + 1
    while True:
        name = _text(ws.cell(r, 1).value)
        if not name:
            break
        stufen = {
            fam: s
            for fam, c in familien_spalten
            if (s := _stufe(ws.cell(r, c).value)) is not None
        }
        blatt.personen.append(
            Person(
                name=" ".join(name.split()),
                team=_text(ws.cell(r, c_team).value) if c_team else None,
                produkte=_text(ws.cell(r, c_produkte).value),
                weitere_bereiche=_text(ws.cell(r, c_weitere).value),
                engpass=_text(ws.cell(r, c_engpass).value) if c_engpass else None,
                validiert_durch=_text(ws.cell(r, c_validiert).value) if c_validiert else None,
                validiert_am=_datum(ws.cell(r, c_datum).value) if c_datum else None,
                notiz=_text(ws.cell(r, c_notiz).value) if c_notiz else None,
                stufen=stufen,
            )
        )
        r += 1

    # Beschreibungen: der Block „Aufgabenfamilie | Enthält aus bestehender Matrix".
    beschreibungen: dict[str, str] = {}
    for zeile in range(r, ws.max_row + 1):
        if _text(ws.cell(zeile, 1).value) == "Aufgabenfamilie":
            for z2 in range(zeile + 1, ws.max_row + 1):
                fam = _text(ws.cell(z2, 1).value)
                if not fam:
                    break
                beschreibungen[fam] = _text(ws.cell(z2, 2).value) or ""
            break

    blatt.familien = [
        Familie(name=fam, beschreibung=beschreibungen.get(fam)) for fam, _ in familien_spalten
    ]
    return blatt


def lies_interviewdatei(daten: bytes, dateiname: str) -> Datei:
    """Alle Bereichsblätter der Interviewdatei; Blätter ohne Kopf entfallen."""
    wb = load_workbook(BytesIO(daten), data_only=True)
    datei = Datei(dateiname=dateiname)
    for ws in wb.worksheets:
        blatt = _lies_blatt(ws)
        if blatt is not None:
            datei.blaetter.append(blatt)
    if not datei.blaetter:
        raise InterviewUnbrauchbar(
            "Keine Bereichsblätter gefunden — fehlt der Mitarbeiter-Kopf?"
        )
    return datei
